#!/usr/bin/env python3
"""
audit_3_5_v2_sampling.py
========================

Generate the Step-3.5 v2 audit workbook for co-author cross-check, per
OSF protocol section 7.4 (recently updated to v0.4 with subagent specification).

Sampling:
    - 100% of promote_to_IN verdicts (every promote reviewed by a co-author).
    - 10% random sample of stay_OUT verdicts (seed = 20260519 to differentiate
      from v1's seed 20260511).

Output: v2_active/audit_3_5_v2.xlsx with three sheets:
    1. instructions
    2. promote_to_IN_review_v2  (review of all v2 promotes)
    3. stay_OUT_sample_v2       (review of ~9 random stay_OUT records)

Each review row contains the screened text plus three columns for the co-author
to fill in: agreement (agree | disagree | uncertain), reviewer_initials, notes.

Usage:
    python3 scripts/v2/audit_3_5_v2_sampling.py
"""

import json
import random
import sys
from pathlib import Path

HERE = Path(__file__).resolve()
STEP_DIR = HERE.parent.parent.parent
sys.path.insert(0, str(STEP_DIR))
import paths_v2 as P  # noqa: E402

try:
    import openpyxl
    from openpyxl.styles import Font, Alignment, PatternFill
    from openpyxl.utils import get_column_letter
except ImportError:
    sys.exit("openpyxl is required: pip3 install --break-system-packages openpyxl")


SEED = 20260519
STAY_OUT_FRACTION = 0.10


def load_jsonl(path):
    if not Path(path).exists():
        return []
    out = []
    with open(path, encoding="utf-8") as f:
        for line in f:
            line = line.strip()
            if not line:
                continue
            try:
                out.append(json.loads(line))
            except json.JSONDecodeError:
                continue
    return out


def main():
    verdicts      = load_jsonl(P.VERDICTS_3_5_PATH)
    input_records = load_jsonl(P.INPUT_PATH)

    if not verdicts:
        sys.exit("No verdicts found; run the sub-screen first.")

    input_by_rid = {r.get("record_id"): r for r in input_records}

    promotes = [v for v in verdicts if v.get("verdict") == "promote_to_IN"]
    stays    = [v for v in verdicts if v.get("verdict") == "stay_OUT"]

    # Reproducible sample
    rng = random.Random(SEED)
    n_sample = int(round(STAY_OUT_FRACTION * len(stays)))
    sample = sorted(stays, key=lambda v: v.get("record_id") or "")
    rng.shuffle(sample)
    sample = sample[:n_sample]

    # Build the workbook
    wb = openpyxl.Workbook()

    # ----- instructions sheet -----
    ws = wb.active
    ws.title = "instructions"

    bold = Font(bold=True)
    fill_header = PatternFill("solid", fgColor="DDDDDD")
    wrap = Alignment(wrap_text=True, vertical="top")

    instructions = [
        "Step 3.5 v2 Audit Workbook",
        "",
        "Project: IEEE TRPMS systematic review on Quantum Computing for Medical Imaging Applications",
        f"Corresponding author: Laszlo Papp",
        f"Sub-screen executed: 2026-05-19",
        f"Total v2 sub-screen verdicts: {len(verdicts)}",
        f"Total promotes (v2): {len(promotes)}",
        f"Total stay_OUT (v2): {len(stays)}",
        "",
        "Sampling parameters (reproducible):",
        f"  Random seed: {SEED}",
        f"  stay_OUT sampling fraction: {STAY_OUT_FRACTION}",
        f"  stay_OUT sample size: {n_sample}",
        "",
        "Cross-check protocol (per OSF section 7.4):",
        "  - Review 100% of v2 promote_to_IN verdicts on the 'promote_to_IN_review_v2' sheet.",
        "  - Review a 10% random sample of stay_OUT verdicts on the 'stay_OUT_sample_v2' sheet.",
        "",
        "For each row, fill in three columns:",
        "  agreement:           one of (agree | disagree | uncertain)",
        "  reviewer_initials:   your initials (e.g. MH for Mathieu Hatt)",
        "  notes:               brief justification or comment",
        "",
        "v2 differs from v1 in the following ways:",
        "  - v2 corpus is the full (non-OA-filtered) set of 6847 deduplicated DOIs.",
        "  - v2 Step-3 added 91 new QC_FOR_IMAGING candidates beyond v1's 88.",
        "  - v2 Step-3.5 sub-screened 96 NEW QUANTUM_SENSING_BIOMED records",
        "    (records already screened in v1 are not re-screened).",
        "",
        "After review, return this workbook to Laszlo Papp; the agreement",
        "rate is recorded as a Transparent Changes entry on the OSF Project page.",
    ]
    for i, text in enumerate(instructions, start=1):
        cell = ws.cell(row=i, column=1, value=text)
        if i == 1:
            cell.font = Font(bold=True, size=14)
    ws.column_dimensions['A'].width = 100

    # ----- promote_to_IN_review_v2 sheet -----
    ws = wb.create_sheet("promote_to_IN_review_v2")
    headers = [
        "record_id", "doi", "ionising_modality_name",
        "ionising_modality_quote", "context_quote", "confidence", "reasoning",
        "title", "abstract", "agreement", "reviewer_initials", "notes",
    ]
    for col, h in enumerate(headers, start=1):
        c = ws.cell(row=1, column=col, value=h)
        c.font = bold
        c.fill = fill_header

    for row_idx, p in enumerate(promotes, start=2):
        rec = input_by_rid.get(p.get("record_id"), {})
        values = [
            p.get("record_id"),
            p.get("doi"),
            p.get("ionising_modality_name"),
            p.get("ionising_modality_quote"),
            p.get("context_quote"),
            p.get("confidence"),
            p.get("reasoning"),
            rec.get("title", ""),
            rec.get("abstract", ""),
            "",  # agreement
            "",  # reviewer_initials
            "",  # notes
        ]
        for col_idx, val in enumerate(values, start=1):
            c = ws.cell(row=row_idx, column=col_idx, value=val)
            c.alignment = wrap

    # Reasonable widths
    widths = [22, 28, 26, 40, 60, 12, 60, 60, 80, 14, 18, 40]
    for col, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(col)].width = w
    ws.freeze_panes = "A2"

    # ----- stay_OUT_sample_v2 sheet -----
    ws = wb.create_sheet("stay_OUT_sample_v2")
    headers = [
        "record_id", "doi", "ionising_modality_present",
        "ionising_modality_quote", "quantum_sensing_in_ionising_context",
        "context_quote", "confidence", "reasoning",
        "title", "abstract",
        "agreement", "reviewer_initials", "notes",
    ]
    for col, h in enumerate(headers, start=1):
        c = ws.cell(row=1, column=col, value=h)
        c.font = bold
        c.fill = fill_header

    for row_idx, s in enumerate(sample, start=2):
        rec = input_by_rid.get(s.get("record_id"), {})
        values = [
            s.get("record_id"),
            s.get("doi"),
            s.get("ionising_modality_present"),
            s.get("ionising_modality_quote", ""),
            s.get("quantum_sensing_in_ionising_context"),
            s.get("context_quote", ""),
            s.get("confidence"),
            s.get("reasoning"),
            rec.get("title", ""),
            rec.get("abstract", ""),
            "",
            "",
            "",
        ]
        for col_idx, val in enumerate(values, start=1):
            c = ws.cell(row=row_idx, column=col_idx, value=val)
            c.alignment = wrap

    widths = [22, 28, 22, 40, 28, 50, 12, 60, 60, 80, 14, 18, 40]
    for col, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(col)].width = w
    ws.freeze_panes = "A2"

    wb.save(P.AUDIT_XLSX)

    print(f"Wrote audit workbook: {P.AUDIT_XLSX}")
    print(f"  promote_to_IN review rows: {len(promotes)}")
    print(f"  stay_OUT sample rows:      {len(sample)} (10% of {len(stays)}, seed={SEED})")


if __name__ == "__main__":
    main()
