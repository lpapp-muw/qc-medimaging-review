# -*- coding: utf-8 -*-
"""build_bibliography.py

Builds references.bib for the 133 included studies.

Design
------
Local first. The Zotero CSL-JSON export and the on-disk DOI metadata cache hold
the metadata the pipeline actually ran on, so they are preferred over a fresh
network fetch. Crossref is queried only for DOIs neither local source covers,
and only when --fetch is given. That is a handful of requests, not 133.

Every record, whatever its source, is cross-checked against the extraction
workbook: the harvested title is compared with the title extracted from the
paper itself. A DOI that resolves to a different article than the one extracted
is the failure mode that silently mis-cites, and it is visible only as a
disagreement between two independent sources.

Citation keys are the DOI with every non-alphanumeric replaced by an
underscore, prefixed 'S'. Corpus keys therefore always begin 'S' followed by a
digit; framing and methods keys are author-year and never do. The mapping is
checked for bijectivity before anything is written.

Usage
-----
  python3 build_bibliography.py                 local sources only
  python3 build_bibliography.py --fetch         also fetch the residue
  python3 build_bibliography.py --root /path    project root, default cwd
  python3 build_bibliography.py --strict        exit non-zero on any gap

Outputs, written beside the workbook unless --out is given:
  references.bib              BibTeX, IEEEtran-compatible
  bibliography_report.txt     coverage, per-field gaps, cross-check failures
  bibliography_missing.csv    DOIs still lacking a required field

Python 3.8 compatible. Requires openpyxl only.
"""
from __future__ import print_function
import os, re, sys, csv, json, glob, time, argparse, unicodedata

try:
    import openpyxl
except ImportError:
    print('FATAL: openpyxl required. Activate the project venv first.')
    sys.exit(2)

EXCLUDED = '10.1109/access.2025.3627877'
REQUIRED = ['author', 'title', 'journal', 'year']
NICE = ['volume', 'number', 'pages']


def doi_key(doi):
    return 'S' + re.sub(r'[^0-9a-zA-Z]', '_', doi.strip().lower())


def norm(s):
    if s is None:
        return ''
    s = unicodedata.normalize('NFKD', str(s))
    s = ''.join(c for c in s if not unicodedata.combining(c))
    s = s.lower()
    s = re.sub(r'&amp;|&lt;|&gt;|&quot;', ' ', s)
    s = re.sub(r'<[^>]+>', ' ', s)
    s = re.sub(r'[^a-z0-9]+', ' ', s)
    return ' '.join(s.split())


def token_overlap(a, b):
    sa, sb = set(norm(a).split()), set(norm(b).split())
    if not sa or not sb:
        return 0.0
    return float(len(sa & sb)) / float(len(sa | sb))


def bib_escape(s):
    if s is None:
        return ''
    s = re.sub(r'&amp;', '&', str(s))
    s = re.sub(r'<[^>]+>', '', s)
    for a, b in [('\\', ''), ('&', r'\&'), ('%', r'\%'), ('$', r'\$'),
                 ('#', r'\#'), ('_', r'\_'), ('{', ''), ('}', ''), ('~', ' ')]:
        s = s.replace(a, b)
    return ' '.join(s.split())


def load_corpus(wb_path):
    wb = openpyxl.load_workbook(wb_path, read_only=True, data_only=True)
    ws = wb['extractions']
    it = ws.iter_rows(min_row=1, values_only=True)
    next(it)
    hdr = list(next(it))
    idx = dict((h, i) for i, h in enumerate(hdr))

    def g(r, n):
        if n not in idx:
            return ''
        v = r[idx[n]]
        return '' if v is None else str(v).strip()

    out = []
    for r in it:
        if not r[0]:
            continue
        doi = g(r, 'doi').strip()
        if doi.lower() == EXCLUDED:
            continue
        out.append({'doi': doi, 'stable_name': g(r, 'stable_name'),
                    'title': g(r, 'title'), 'journal': g(r, 'journal'),
                    'authors': g(r, 'authors')})
    return out


def load_years(human_path):
    years = {}
    if not os.path.exists(human_path):
        return years
    wb = openpyxl.load_workbook(human_path, data_only=True)
    ws = wb['extractions']
    for r in ws.iter_rows(min_row=3, values_only=True):
        if r[0] and len(r) > 4 and r[4] not in (None, ''):
            try:
                years[str(r[0]).strip().lower()] = int(r[4])
            except (ValueError, TypeError):
                pass
    return years


def csl_records(root):
    out = {}
    pat = os.path.join(root, '3-AbstractRetrieve', 'inputs', 'zotero_csl_json_v1', '*.json')
    for f in sorted(glob.glob(pat)):
        try:
            data = json.load(open(f))
        except Exception as e:
            print('  warning: cannot parse {}: {}'.format(os.path.basename(f), e))
            continue
        if isinstance(data, dict):
            data = [data]
        for rec in data:
            d = (rec.get('DOI') or '').strip().lower()
            if d:
                out.setdefault(d, rec)
    return out


def cache_records(root):
    """DOI metadata cache. Format auto-detected: any JSON object carrying a DOI."""
    out = {}
    base = os.path.join(root, '.doi_meta_cache_v3')
    if not os.path.isdir(base):
        return out
    for f in glob.glob(os.path.join(base, '**', '*.json'), recursive=True):
        try:
            data = json.load(open(f))
        except Exception:
            continue
        cands = data if isinstance(data, list) else [data]
        for rec in cands:
            if not isinstance(rec, dict):
                continue
            if 'message' in rec and isinstance(rec['message'], dict):
                rec = rec['message']
            d = (rec.get('DOI') or rec.get('doi') or '').strip().lower()
            if d:
                out.setdefault(d, rec)
    return out


def crossref_fetch(doi, mailto):
    import urllib.request
    url = 'https://api.crossref.org/works/{}?mailto={}'.format(doi, mailto)
    req = urllib.request.Request(
        url, headers={'User-Agent': 'IEEE-TRPMS-review/1.0 (mailto:{})'.format(mailto)})
    fh = urllib.request.urlopen(req, timeout=30)
    try:
        return json.load(fh).get('message', {})
    finally:
        fh.close()


def authors_from(rec):
    a = rec.get('author')
    if not a:
        return ''
    parts = []
    for x in a:
        fam = (x.get('family') or '').strip()
        giv = (x.get('given') or '').strip()
        if fam and giv:
            parts.append('{}, {}'.format(fam, giv))
        elif fam:
            parts.append(fam)
        elif x.get('literal'):
            parts.append(x['literal'].strip())
    return ' and '.join(parts)


def year_from(rec):
    for k in ('issued', 'published-print', 'published-online', 'published'):
        v = rec.get(k)
        if isinstance(v, dict) and v.get('date-parts'):
            dp = v['date-parts'][0]
            if dp and dp[0]:
                try:
                    return int(dp[0])
                except (ValueError, TypeError):
                    pass
    return None


def first(v):
    if isinstance(v, list):
        return v[0] if v else ''
    return v or ''


def to_entry(rec):
    return {'author': authors_from(rec),
            'title': first(rec.get('title')),
            'journal': first(rec.get('container-title')) or first(rec.get('journalAbbreviation')),
            'year': year_from(rec),
            'volume': (rec.get('volume') or '').strip(),
            'number': (rec.get('issue') or '').strip(),
            'pages': (rec.get('page') or '').strip()}


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--root', default=os.getcwd())
    ap.add_argument('--workbook', default=None)
    ap.add_argument('--human', default=None)
    ap.add_argument('--out', default=None)
    ap.add_argument('--fetch', action='store_true')
    ap.add_argument('--mailto', default='laszlo.papp@meduniwien.ac.at')
    ap.add_argument('--title-threshold', type=float, default=0.60)
    ap.add_argument('--strict', action='store_true')
    args = ap.parse_args()

    root = os.path.abspath(args.root)
    wb_path = args.workbook or os.path.join(
        root, '7-Manuscript', 'supp_enrich', 'out', 'extractions_ai_enriched_v134.xlsx')
    human_path = args.human or os.path.join(root, '6-InterRaterKappa', 'human.xlsx')
    outdir = args.out or os.path.dirname(wb_path)

    if not os.path.exists(wb_path):
        print('FATAL: workbook not found at {}'.format(wb_path))
        sys.exit(2)

    print('root     : {}'.format(root))
    print('workbook : {}'.format(wb_path))
    print('')

    corpus = load_corpus(wb_path)
    years = load_years(human_path)
    csl = csl_records(root)
    cache = cache_records(root)
    print('included studies      : {}'.format(len(corpus)))
    print('years from human.xlsx : {}'.format(len(years)))
    print('CSL-JSON records      : {}'.format(len(csl)))
    print('DOI cache records     : {}'.format(len(cache)))
    print('')

    keys = {}
    for c in corpus:
        k = doi_key(c['doi'])
        if k in keys:
            print('FATAL: key collision {} for {} and {}'.format(k, keys[k], c['doi']))
            sys.exit(2)
        keys[k] = c['doi']
    print('citation keys         : {} distinct, bijection OK'.format(len(keys)))

    entries, fetchfail, missing, mismatch, wb_fallback = [], [], [], [], []
    src_count = {'csl': 0, 'cache': 0, 'crossref': 0, 'none': 0}

    for c in corpus:
        d = c['doi'].lower()
        rec, source = None, 'none'
        if d in csl:
            rec, source = csl[d], 'csl'
        elif d in cache:
            rec, source = cache[d], 'cache'
        elif args.fetch:
            try:
                rec, source = crossref_fetch(c['doi'], args.mailto), 'crossref'
                time.sleep(1.0)
            except Exception as e:
                fetchfail.append('{}: {}'.format(c['doi'], e))
        e = to_entry(rec) if rec else {}
        src_count[source] += 1

        if not e.get('year') and d in years:
            e['year'] = years[d]

        wb_used = []
        if not e.get('author') and c.get('authors'):
            names = [x.strip() for x in c['authors'].split(';') if x.strip()]
            if names:
                e['author'] = ' and '.join(names)
                wb_used.append('author')
        if not e.get('title') and c.get('title'):
            e['title'] = c['title']; wb_used.append('title')
        if not e.get('journal') and c.get('journal'):
            e['journal'] = c['journal']; wb_used.append('journal')
        if wb_used:
            wb_fallback.append((c['doi'], ';'.join(wb_used)))

        if rec:
            ov = token_overlap(e.get('title'), c['title'])
            if ov < args.title_threshold:
                mismatch.append((c['doi'], source, round(ov, 2),
                                 (c['title'] or '')[:70], (e.get('title') or '')[:70]))

        gaps = [f for f in REQUIRED if not e.get(f)]
        if gaps:
            missing.append((c['doi'], source, ';'.join(gaps),
                            ';'.join(f for f in NICE if not e.get(f))))

        e['key'] = doi_key(c['doi'])
        e['doi'] = c['doi']
        e['source'] = source
        entries.append(e)

    bib_path = os.path.join(outdir, 'references.bib')
    fh = open(bib_path, 'w')
    fh.write('% Included studies, IEEE TRPMS systematic review.\n')
    fh.write('% Keys are the DOI with non-alphanumerics underscored, prefixed S.\n')
    fh.write('% Generated by build_bibliography.py. Do not hand-edit keys.\n\n')
    for e in sorted(entries, key=lambda x: x['key']):
        fh.write('@article{%s,\n' % e['key'])
        fh.write('  author  = {%s},\n' % bib_escape(e.get('author')))
        fh.write('  title   = {%s},\n' % bib_escape(e.get('title')))
        fh.write('  journal = {%s},\n' % bib_escape(e.get('journal')))
        if e.get('year'):
            fh.write('  year    = {%s},\n' % e['year'])
        for f in ('volume', 'number', 'pages'):
            if e.get(f):
                fh.write('  %-7s = {%s},\n' % (f, bib_escape(e[f])))
        fh.write('  doi     = {%s},\n' % e['doi'])
        fh.write('}\n\n')
    fh.close()

    rep_path = os.path.join(outdir, 'bibliography_report.txt')
    rf = open(rep_path, 'w')

    def w(s=''):
        rf.write(s + '\n')
        print(s)

    w('')
    w('=' * 78)
    w('BIBLIOGRAPHY REPORT')
    w('=' * 78)
    w('  entries written      : {}'.format(len(entries)))
    w('  source CSL-JSON      : {}'.format(src_count['csl']))
    w('  source DOI cache     : {}'.format(src_count['cache']))
    w('  source Crossref      : {}'.format(src_count['crossref']))
    w('  no metadata found    : {}'.format(src_count['none']))
    w('')
    complete = sum(1 for e in entries if all(e.get(f) for f in REQUIRED))
    w('  complete on author/title/journal/year: {} of {}'.format(complete, len(entries)))
    for f in REQUIRED + NICE:
        n = sum(1 for e in entries if e.get(f))
        w('    {:<8} present in {:>3} of {}'.format(f, n, len(entries)))
    w('')
    w('  WORKBOOK FALLBACK, fields filled from the quote-grounded extraction')
    w('  studies affected: {}'.format(len(wb_fallback)))
    for doi, fields in wb_fallback:
        w('    {}  ->  {}'.format(doi, fields))
    if not wb_fallback:
        w('    none')
    w('')
    w('  CROSS-CHECK against the extraction workbook')
    w('  title overlap below {:.2f}: {}'.format(args.title_threshold, len(mismatch)))
    for doi, s, ov, wt, ft in mismatch:
        w('    {} [{}] overlap {}'.format(doi, s, ov))
        w('        workbook : {}'.format(wt))
        w('        harvested: {}'.format(ft))
    if not mismatch:
        w('    none; every harvested title matches the extracted title')
    w('')
    if fetchfail:
        w('  FETCH FAILURES')
        for r in fetchfail:
            w('    ' + r)
        w('')
    w('  Re-run with --fetch to query Crossref for DOIs neither local source')
    w('  covers, at one request per second.')
    rf.close()

    mf = open(os.path.join(outdir, 'bibliography_missing.csv'), 'w')
    wr = csv.writer(mf)
    wr.writerow(['doi', 'source', 'missing_required', 'missing_optional'])
    for row in missing:
        wr.writerow(row)
    mf.close()

    print('')
    print('  written: {}'.format(bib_path))
    print('           {}'.format(rep_path))
    print('           {}'.format(os.path.join(outdir, 'bibliography_missing.csv')))
    print('  studies still missing a required field: {}'.format(len(missing)))

    if args.strict and (missing or mismatch):
        sys.exit(1)


if __name__ == '__main__':
    main()
