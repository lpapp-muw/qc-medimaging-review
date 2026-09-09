"""
check_missing.py — Cross-check on-disk PDFs against the 191-candidate DOI set.

Reports which candidate DOIs have no PDF in PDFs_step4/, with their Stage and
Confidence, for the PRISMA "full text not retrievable" exclusion list.

Run from extract_v1/. Python 3.8.
"""

import sys
from pathlib import Path
import paths_step4 as P
from openpyxl import load_workbook


def doi_to_pdf_stable(doi):
    # consolidation convention: '/' -> '_', everything else preserved
    return doi.replace("/", "_")


def main():
    # 1. Candidate DOIs + metadata
    wb = load_workbook(str(P.CANDIDATES_XLSX), data_only=True, read_only=True)
    ws = wb["all_candidates"]
    rows = ws.iter_rows(values_only=True)
    hdr = [str(h).strip() if h else "" for h in next(rows)]
    di = hdr.index("DOI")
    si = hdr.index("Stage")
    ci = hdr.index("Confidence")
    cand = {}
    for r in rows:
        if r and di < len(r) and r[di]:
            cand[str(r[di]).strip().lower()] = (r[si] if si < len(r) else "",
                                                 r[ci] if ci < len(r) else "")
    wb.close()

    # 2. On-disk PDF stems
    on_disk = set()
    for p in P.PDFS_STEP4.glob("*.pdf"):
        on_disk.add(p.stem.lower())

    # 3. Match each candidate by its expected stable filename
    missing = []
    present = 0
    for doi, (stage, conf) in cand.items():
        stable = doi_to_pdf_stable(doi).lower()
        if stable in on_disk:
            present += 1
        else:
            missing.append((doi, stage, conf))

    print("Candidates           :", len(cand))
    print("PDFs on disk          :", len(on_disk))
    print("Matched (present)     :", present)
    print("Missing               :", len(missing))
    print()
    if missing:
        print("MISSING DOIs (PRISMA 'full text not retrievable'):")
        print("{:<42s} {:<18s} {}".format("DOI", "Stage", "Confidence"))
        for doi, stage, conf in sorted(missing):
            print("{:<42s} {:<18s} {}".format(doi, str(stage), str(conf)))

    # 4. Reverse check: any on-disk PDF NOT in candidate set? (stray files)
    cand_stables = set(doi_to_pdf_stable(d).lower() for d in cand)
    stray = sorted(s for s in on_disk if s not in cand_stables)
    if stray:
        print()
        print("STRAY PDFs on disk not matching any candidate DOI:")
        for s in stray:
            print(" ", s)

    return 0


if __name__ == "__main__":
    sys.exit(main())
