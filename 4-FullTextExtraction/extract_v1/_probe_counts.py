import openpyxl
from collections import Counter
ws = openpyxl.load_workbook('extractions_ai.xlsx', data_only=True)['extractions']
hdr = [c.value for c in ws[2]]; col = {h: i for i, h in enumerate(hdr) if h}
k = 'paradigm_subcategory (derived)'
vals = [row[col[k]] for row in ws.iter_rows(min_row=3, values_only=True) if row[col['doi']]]
filled = sum(1 for v in vals if v not in (None, ''))
print("rows:", len(vals), "| filled:", filled, "| bare other_unresolved:", sum(1 for v in vals if v == 'other_unresolved'))
for v, n in Counter(vals).most_common():
    print(f"  {n:4d}  {v}")
