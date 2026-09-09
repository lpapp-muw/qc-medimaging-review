import openpyxl
ws = openpyxl.load_workbook('extractions_ai.xlsx', data_only=True)['extractions']
hdr = [c.value for c in ws[2]]
col = {h: i for i, h in enumerate(hdr) if h}
want = ['doi', 'paradigm', 'paradigm_subcategory (derived)',
        'paradigm_facets (derived)', 'paradigm_resolution_basis (derived)']
missing = [w for w in want if w not in col]
if missing:
    print("MISSING COLUMNS:", missing)
    raise SystemExit(1)
print(' || '.join(want))
probe = {'10.1038/s41598-025-08453-w', '10.1088/1748-0221/18/07/p07007', '10.3934/math.2025499'}
for row in ws.iter_rows(min_row=3, values_only=True):
    d = row[col['doi']]
    if d and str(d).lower() in probe:
        print(' || '.join(str(row[col[w]]) for w in want))
