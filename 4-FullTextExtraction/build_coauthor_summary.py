#!/usr/bin/env python3
"""
build_coauthor_summary.py
=========================

Generate Step4_CoAuthor_Summary.docx for distribution to co-authors,
summarizing:

1. Project context and PRISMA flow numbers (v1 + v2).
2. Step-4 candidate set composition (191 records).
3. PDF retrieval status (89 have_local + 16 retrieved_OA + N missing).
4. Workflow for co-author full-text review.
5. Embedded table of the still-missing PDFs, sorted by publisher,
   so co-authors can sort by their institution's subscriptions.

Reads:
    step4_candidates_v2.xlsx (sheets: all_candidates, missing_pdfs)

Writes:
    Step4_CoAuthor_Summary.docx

Dependencies:
    python-docx (pip3 install --break-system-packages python-docx)
    openpyxl

Usage:
    python3 4-FullTextExtraction/build_coauthor_summary.py
"""

import sys
from collections import Counter
from datetime import date
from pathlib import Path

try:
    import openpyxl
except ImportError:
    sys.exit("openpyxl required: pip3 install --break-system-packages openpyxl")

try:
    from docx import Document
    from docx.shared import Pt, Cm, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
except ImportError:
    sys.exit("python-docx required: pip3 install --break-system-packages python-docx")


PROJECT_ROOT = Path("/home/lpapp/IEEE_SYS_REV")
STEP4_DIR    = PROJECT_ROOT / "4-FullTextExtraction"
WORKBOOK     = STEP4_DIR / "step4_candidates_v2.xlsx"
PDF_FOLDER   = STEP4_DIR / "PDFs_step4"
OUT_DOCX     = STEP4_DIR / "Step4_CoAuthor_Summary.docx"


def set_arial(run, size_pt=11, bold=False):
    run.font.name = "Arial"
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:ascii"), "Arial")
    rFonts.set(qn("w:hAnsi"), "Arial")
    rFonts.set(qn("w:cs"), "Arial")
    rFonts.set(qn("w:eastAsia"), "Arial")


def add_para(doc, text, *, bold=False, size=11, align=None, before=0, after=120):
    p = doc.add_paragraph()
    if align == "center":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    run = p.add_run(text)
    set_arial(run, size, bold=bold)
    return p


def add_heading(doc, text, level=1):
    size_map = {1: 16, 2: 13, 3: 11}
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(text)
    set_arial(run, size_map.get(level, 11), bold=True)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text)
    set_arial(run, 11)
    return p


def add_table(doc, headers, rows, col_widths_cm):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Light Grid Accent 1"
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    # Header
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = ""
        p = hdr[i].paragraphs[0]
        run = p.add_run(h)
        set_arial(run, 10, bold=True)
        # Set width
        if col_widths_cm and i < len(col_widths_cm):
            hdr[i].width = Cm(col_widths_cm[i])
    # Body rows
    for row_data in rows:
        cells = table.add_row().cells
        for i, val in enumerate(row_data):
            cells[i].text = ""
            p = cells[i].paragraphs[0]
            run = p.add_run(str(val) if val is not None else "")
            set_arial(run, 9)
            if col_widths_cm and i < len(col_widths_cm):
                cells[i].width = Cm(col_widths_cm[i])
    return table


def doi_publisher_prefix(doi):
    if not doi or not doi.lower().startswith("10."):
        return ""
    return doi.lower().split("/", 1)[0]


def main():
    if not WORKBOOK.exists():
        sys.exit(f"Missing workbook: {WORKBOOK}")

    wb = openpyxl.load_workbook(WORKBOOK)
    all_ws = wb["all_candidates"]
    miss_ws = wb["missing_pdfs"]

    all_headers = [c.value for c in all_ws[1]]
    miss_headers = [c.value for c in miss_ws[1]]

    # Index columns
    a_doi    = all_headers.index("DOI") + 1
    a_status = all_headers.index("PDF_status") + 1
    a_stage  = all_headers.index("Stage") + 1
    a_conf   = all_headers.index("Confidence") + 1
    a_title  = all_headers.index("Title") + 1
    a_journal = all_headers.index("Journal") + 1
    a_pub    = all_headers.index("Publisher_guess") + 1
    a_year   = all_headers.index("Year") + 1

    # Walk all_candidates
    total = 0
    have_local = 0
    retrieved_OA = 0
    missing = 0
    by_stage = Counter()
    by_conf = Counter()
    missing_rows = []  # for the missing-PDFs table later

    for row in all_ws.iter_rows(min_row=2, values_only=True):
        if not row[a_doi - 1]:
            continue
        total += 1
        status = row[a_status - 1]
        stage  = row[a_stage - 1]
        conf   = row[a_conf - 1]
        by_stage[stage] += 1
        by_conf[conf] += 1
        if status == "have_local":
            have_local += 1
        elif status == "retrieved_OA":
            retrieved_OA += 1
        else:
            missing += 1
            missing_rows.append({
                "doi":    row[a_doi - 1],
                "title":  row[a_title - 1],
                "year":   row[a_year - 1],
                "journal": row[a_journal - 1],
                "pub":    row[a_pub - 1],
                "stage":  stage,
                "conf":   conf,
            })

    # PDF folder count
    pdfs_in_folder = len(list(PDF_FOLDER.glob("*.pdf"))) if PDF_FOLDER.exists() else 0

    # Build docx
    doc = Document()
    # Set page margins to ~2cm
    for section in doc.sections:
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin = Cm(2)
        section.right_margin = Cm(2)

    # Title
    add_para(doc, "Step 4 Co-Author Summary",
             bold=True, size=18, align="center", after=12)
    add_para(doc, "Quantum Computing for Medical Imaging Applications "
             "— invited systematic review, IEEE TRPMS",
             size=11, align="center", after=4)
    add_para(doc, f"Corresponding author: Laszlo Papp",
             size=10, align="center", after=4)
    add_para(doc, f"Generated: {date.today().isoformat()}",
             size=10, align="center", after=18)

    # 1. Overview
    add_heading(doc, "1. Overview", 1)
    add_para(doc,
        "This document summarises the state of the systematic-review "
        "pipeline at the start of Step 4 (full-text review). The pipeline "
        "ran in two passes: a v1 pilot on the open-access subset of the "
        "literature to lock the 10-category screening taxonomy, and a v2 "
        "full-corpus pass that applied the locked taxonomy to all 6847 "
        "deduplicated DOIs returned by PubMed, Web of Science, and IEEE "
        "Xplore on 30 January 2026.")
    add_para(doc,
        "The OSF protocol v0.6 documents the methodology and is "
        "referenced for all details. This document is a working "
        "companion for the co-author review effort.")

    # 2. PRISMA flow
    add_heading(doc, "2. PRISMA flow (v2 corpus is canonical)", 1)
    add_bullet(doc, "Database hits (30 Jan 2026): PubMed 1569, WoS 2201, IEEE Xplore 123, plus v2 non-OA expansion. Combined raw: 8961 entries.")
    add_bullet(doc, "After DOI deduplication: 6847 unique DOIs (v2 canonical corpus).")
    add_bullet(doc, "Excluded prior to Stage 3: 8 records with no retrievable metadata from Crossref / OpenAlex / EuropePMC (PRISMA reason: \"no metadata retrievable\").")
    add_bullet(doc, "Records screened by AI-assisted abstract classification (Stage 3): 6839.")
    add_bullet(doc, f"In-scope after Stage 3: 179 QC_FOR_IMAGING (120 high-confidence, 59 low-confidence).")
    add_bullet(doc, f"In-scope after Stage 3.5 ionising sub-screen: +12 promotes from QUANTUM_SENSING_BIOMED (4 from v1, 8 from v2; all high-confidence; covering PET, radiation-therapy dosimetric imaging, proton radiography, planar X-ray, CT).")
    add_bullet(doc, f"Total Step-4 candidates: {total}.")

    # 3. Step-4 candidate set
    add_heading(doc, "3. Step-4 candidate set composition", 1)
    headers = ["Stage", "Count"]
    rows = [(s, n) for s, n in by_stage.most_common()]
    rows.append(("Total", total))
    add_table(doc, headers, rows, [6.5, 3])

    add_para(doc, "Confidence distribution:", bold=True, after=4)
    rows = [(c, n) for c, n in by_conf.most_common()]
    add_table(doc, ["Confidence", "Count"], rows, [6.5, 3])

    # 4. PDF status
    add_heading(doc, "4. PDF retrieval status", 1)
    add_bullet(doc, f"Have local PDF (from v1 Zotero storage): {have_local}")
    add_bullet(doc, f"Retrieved via open-access APIs (v2 retrieval pass): {retrieved_OA}")
    add_bullet(doc, f"Still missing (need institutional access): {missing}")
    add_bullet(doc, f"Total PDFs consolidated in PDFs_step4/ folder: {pdfs_in_folder}")
    add_para(doc,
        "The {} PDFs we have are gathered in the folder ".format(pdfs_in_folder) +
        "PDFs_step4/ (filename convention: <doi-with-slash-replaced>.pdf). "
        "The {} remaining papers are listed in section 6 below; the same "
        "list with all columns is on sheet 'missing_pdfs' of the "
        "accompanying step4_candidates_v2.xlsx workbook.".format(missing))

    # 5. Workflow
    add_heading(doc, "5. Workflow for co-author full-text review", 1)
    add_bullet(doc,
        "Open step4_candidates_v2.xlsx, sheet 'all_candidates'. Filter by "
        "the 'Assigned_reviewer' column to claim records.")
    add_bullet(doc,
        "For each claimed record, open the PDF from PDFs_step4/ "
        "(when 'PDF_status' is have_local or retrieved_OA).")
    add_bullet(doc,
        "Apply the Step-4 eligibility criteria from the OSF protocol "
        "(section 7.5). Fill the 'Decision' column with one of: "
        "include / exclude / uncertain.")
    add_bullet(doc,
        "For low-confidence Stage-3 records (Confidence column = low), "
        "expect a higher exclusion rate; the AI screen was uncertain.")
    add_bullet(doc,
        "For missing PDFs (section 6 of this document), attempt retrieval "
        "via your institutional access. Record your initials in the "
        "'PDF_retrieved_by' column and drop the file in PDFs_step4/ with "
        "the same naming convention.")
    add_bullet(doc,
        "Use the 'Notes' column for any reasoning that should appear in "
        "the final manuscript or the PRISMA exclusion table.")

    # 6. Missing-PDFs table
    add_heading(doc, "6. Missing PDFs — please attempt institutional retrieval", 1)
    add_para(doc,
        "The following {} papers are in the Step-4 candidate set but no "
        "PDF was retrievable via open-access channels. Each co-author "
        "should review the table below, identify papers their institution "
        "is likely to have subscriptions for (using the Publisher column), "
        "attempt retrieval, and drop the PDFs into PDFs_step4/ with the "
        "DOI-derived filename convention.".format(missing))

    # Build the table sorted by publisher then DOI
    missing_rows.sort(key=lambda r: (str(r["pub"]) or "", r["doi"]))
    headers = ["DOI", "Title", "Year", "Journal", "Publisher"]
    rows = []
    for r in missing_rows:
        title = r["title"] or ""
        if len(title) > 90:
            title = title[:87] + "..."
        rows.append([
            r["doi"],
            title,
            r["year"] or "",
            r["journal"] or "",
            r["pub"] or "",
        ])
    add_table(doc, headers, rows, [4.5, 6.5, 1, 3, 2.5])

    # Publisher summary
    add_heading(doc, "7. Publisher summary for the missing set", 1)
    add_para(doc,
        "Distribution of the {} missing papers by publisher (for "
        "institutional-access planning):".format(missing))
    by_pub = Counter(r["pub"] for r in missing_rows)
    rows = [(p or "(unknown)", n) for p, n in by_pub.most_common()]
    add_table(doc, ["Publisher", "Missing count"], rows, [10, 3])

    # Footer
    add_para(doc, "", after=12)
    add_para(doc,
        "Companion files in this distribution: "
        "step4_candidates_v2.xlsx (full candidate workbook with assignment "
        "and decision columns), OSF_Protocol_QC_MedImaging_v0_6.docx "
        "(systematic review protocol; OSF deposit candidate), and "
        "PDFs_step4/ (folder of 105 retrieved PDFs).",
        size=10)

    doc.save(OUT_DOCX)
    print(f"Wrote: {OUT_DOCX}")
    print(f"  Total candidates:    {total}")
    print(f"  have_local:          {have_local}")
    print(f"  retrieved_OA:        {retrieved_OA}")
    print(f"  missing:             {missing}")
    print(f"  PDFs in folder:      {pdfs_in_folder}")


if __name__ == "__main__":
    main()
