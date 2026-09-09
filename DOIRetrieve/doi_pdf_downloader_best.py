#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
doi_pdf_downloader_best.py  (Python 3.8+)

Input
-----
XLSX file with a column header "DOI" (each row has one DOI).

Output
------
1) PDFs are downloaded into ./PDFs (relative to this script) unless --pdf-dir is provided.
2) The XLSX file is updated in-place (or to --output) with a column header "PDF":
      - if download succeeded: stored filename (doi-based stable name)
      - else: "Undetected"

Legal scope
-----------
This script only uses lawful sources:
- Open-access discovery APIs: OpenAlex, Semantic Scholar, Crossref, Europe PMC, NCBI PMC OA
- Optional: Unpaywall (requires an email as per Unpaywall policy)
- DOI resolver (doi.org) and publisher landing pages
- Publisher URL patterns that may work for open access or when you have legitimate access

It does NOT use Sci-Hub/LibGen or any piracy source.

Why you may see "status=200 but not PDF"
----------------------------------------
Some publishers (notably ScienceDirect) return an HTML "Preparing your download" page with 200 OK,
not the PDF. This script detects that, scrapes embedded PDF URLs (Elsevier CDN, ScienceDirect assets,
etc.), and retries automatically.

Why you may see 403
-------------------
403 usually means paywall / anti-bot / missing session cookies. This script will:
- retry with cookies warmed up by visiting the referer/landing page
- try alternative OA endpoints (Elsevier CDN, OA repos)
If still blocked, it marks "Undetected".

Dependencies
------------
Required:
  python -m pip install requests openpyxl

Optional (validation):
  python -m pip install pypdf
  python -m pip install pdf2doi     # optional; may require upgrading pdfminer.six

Optional (metadata helper):
  JabRef installed + --jabref PATH or env JABREF_PATH.

Usage
-----
  python doi_pdf_downloader_best.py --input my_dois.xlsx

Recommended (improves OA hit rate):
  python doi_pdf_downloader_best.py --input my_dois.xlsx --unpaywall-email you@domain.com

Politeness throttling (avoid blocks):
  python doi_pdf_downloader_best.py --input my_dois.xlsx --sleep 1.2 --request-sleep 0.2

Create a sample XLSX:
  python doi_pdf_downloader_best.py --make-sample sample_dois.xlsx
"""

from __future__ import annotations

import argparse
import hashlib
import json
import os
import random
import re
import shutil
import subprocess
import sys
import time
import xml.etree.ElementTree as ET
from dataclasses import dataclass
from pathlib import Path
from typing import Dict, Iterable, List, Optional, Sequence, Set, Tuple
from urllib.parse import quote, urljoin, urlparse

# -----------------------------
# Dependency helpers
# -----------------------------


def _pip_install(packages: Sequence[str]) -> bool:
    try:
        cmd = [sys.executable, "-m", "pip", "install", "--upgrade"] + list(packages)
        print(f"[deps] Installing: {' '.join(packages)}")
        subprocess.check_call(cmd)
        return True
    except Exception as e:
        print(f"[deps] pip install failed: {e}", file=sys.stderr)
        return False


def ensure_import(module_name: str, pip_name: Optional[str] = None, auto_install: bool = False) -> bool:
    try:
        __import__(module_name)
        return True
    except ImportError:
        if not auto_install:
            return False
        pip_pkg = pip_name or module_name
        return _pip_install([pip_pkg]) and ensure_import(module_name, pip_name, auto_install=False)


# -----------------------------
# Constants / regex
# -----------------------------

DOI_COL = "DOI"
PDF_COL = "PDF"

DEFAULT_TIMEOUT_S = 45
DEFAULT_SLEEP_S = 1.0            # between DOIs
DEFAULT_REQ_SLEEP_S = 0.15       # between URL attempts
DEFAULT_JITTER_S = 0.25          # random jitter on top of sleeps
MAX_PDF_MB = 250
MAX_HTML_SNIFF_BYTES = 2 * 1024 * 1024

DEFAULT_UA = (
    "Mozilla/5.0 (X11; Linux x86_64) "
    "AppleWebKit/537.36 (KHTML, like Gecko) "
    "Chrome/122.0.0.0 Safari/537.36"
)

DOI_PATTERN = re.compile(r"(10\.\d{4,9}/[-._;()/:A-Z0-9]+)", re.IGNORECASE)

# Elsevier / ScienceDirect patterns
ELSEVIER_PII_RE = re.compile(r"(?:/pii/|/retrieve/pii/)([A-Za-z0-9]+)", re.IGNORECASE)

# URL finder (absolute)
ABS_URL_RE = re.compile(r"https?://[^\s\"\'<>]+", flags=re.IGNORECASE)
# URLs that start with // (scheme-relative)
DOUBLESLASH_URL_RE = re.compile(r"//[^\s\"\'<>]+", flags=re.IGNORECASE)
# Relative URL-like strings in quotes that look PDF-ish (often appears inside JS)
REL_PDFISH_RE = re.compile(
    r"""["\'](/[^"\']*(?:\.pdf|/pdf|/pdfft|pdfdirect|epdf|downloadpdf|pdfdownload)[^"\']*)["\']""",
    flags=re.IGNORECASE,
)
# href finder (relative or absolute)
HREF_RE = re.compile(r"""href\s*=\s*["\']([^"\']+)["\']""", flags=re.IGNORECASE)
META_CITATION_PDF_RE = re.compile(r"""citation_pdf_url[^>]*content=["\']([^"\']+)["\']""", flags=re.IGNORECASE)
META_PDF_RE = re.compile(r"""name=["\']citation_pdf_url["\'][^>]*content=["\']([^"\']+)["\']""", flags=re.IGNORECASE)


# -----------------------------
# Small utils
# -----------------------------


def normalize_doi(raw: str) -> str:
    s = (raw or "").strip()
    if not s:
        return ""
    s = s.replace("\u200b", "").strip()
    s = re.sub(r"^\s*(https?://)?(dx\.)?doi\.org/", "", s, flags=re.IGNORECASE)
    s = re.sub(r"^\s*doi\s*:\s*", "", s, flags=re.IGNORECASE)
    m = DOI_PATTERN.search(s)
    if m:
        return m.group(1).strip()
    return s.strip()


def safe_filename_from_doi(doi: str) -> str:
    d = normalize_doi(doi)
    base = re.sub(r"[^A-Za-z0-9._-]+", "_", d).strip("_")
    if not base:
        base = "unknown_doi"
    if len(base) > 160:
        h = hashlib.sha256(d.encode("utf-8")).hexdigest()[:16]
        base = base[:140] + "_" + h
    return base + ".pdf"


def ensure_dir(path: Path) -> None:
    path.mkdir(parents=True, exist_ok=True)


def sleep_polite(base_s: float, jitter_s: float) -> None:
    if base_s <= 0 and jitter_s <= 0:
        return
    time.sleep(max(0.0, base_s) + (random.random() * max(0.0, jitter_s)))


def pdf_magic_in(b: bytes) -> bool:
    """
    Robust PDF detection: accept if '%PDF' appears within first 1024 bytes.
    (Some servers prepend a few bytes.)
    """
    if not b:
        return False
    idx = b.find(b"%PDF")
    return (idx != -1) and (idx < 1024)


def is_pdf_file(path: Path) -> bool:
    try:
        with path.open("rb") as f:
            head = f.read(2048)
        return pdf_magic_in(head)
    except Exception:
        return False


def looks_like_pdf_url(url: str) -> bool:
    u = (url or "").lower()
    if u.endswith(".pdf") or ".pdf?" in u:
        return True
    # Common endpoints without .pdf
    tokens = [
        "/pdf",
        "/pdfft",
        "pdfdirect",
        "/epdf",
        "downloadpdf",
        "pdfdownload",
        "reader.elsevier.com/reader",
        "pdf.sciencedirectassets.com",
        "ars.els-cdn.com",
        "pmc.ncbi.nlm.nih.gov/articles/",
        "europepmc",
    ]
    return any(t in u for t in tokens)


def dedupe_candidates(cands: Iterable["Candidate"]) -> List["Candidate"]:
    seen: Set[str] = set()
    out: List[Candidate] = []
    for c in cands:
        u = (c.url or "").strip()
        if not u:
            continue
        u = u.split("#", 1)[0]
        if u in seen:
            continue
        seen.add(u)
        out.append(Candidate(url=u, source=c.source, referer=c.referer))
    return out


def candidate_priority(c: "Candidate") -> int:
    """
    Lower score tried earlier.
    """
    u = c.url.lower()
    s = 100

    # Direct PDFs first
    if u.endswith(".pdf") or ".pdf?" in u:
        s -= 30

    # Elsevier OA direct links
    if "ars.els-cdn.com" in u:
        s -= 45
    if "pdf.sciencedirectassets.com" in u:
        s -= 25
    if "reader.elsevier.com" in u:
        s -= 20

    # OA APIs
    if c.source.startswith("unpaywall"):
        s -= 25
    if c.source.startswith("openalex"):
        s -= 15
    if c.source.startswith("semanticscholar"):
        s -= 15
    if c.source.startswith("europepmc") or c.source.startswith("pmc"):
        s -= 12
    if c.source.startswith("crossref"):
        s -= 8

    # Publisher patterns
    if "sciencedirect" in c.source or "elsevier" in c.source:
        s -= 6
    if "wiley" in c.source:
        s -= 4
    if "springer" in c.source:
        s -= 3
    if "acs" in c.source:
        s -= 3
    if "ieee" in c.source:
        s -= 2

    # HTML scrape followups late-ish but before random landings
    if c.source.startswith("html:"):
        s -= 5

    return s


# -----------------------------
# HTTP session
# -----------------------------


def make_session() -> "requests.Session":
    import requests  # type: ignore

    s = requests.Session()
    s.headers.update(
        {
            "User-Agent": DEFAULT_UA,
            "Accept-Language": "en-US,en;q=0.9",
            "Accept": "application/pdf,application/xhtml+xml,text/html;q=0.9,*/*;q=0.8",
            "DNT": "1",
            "Upgrade-Insecure-Requests": "1",
        }
    )

    # Retries for transient server issues
    try:
        from requests.adapters import HTTPAdapter  # type: ignore
        from urllib3.util.retry import Retry  # type: ignore

        try:
            retry = Retry(
                total=4,
                connect=4,
                read=4,
                backoff_factor=0.7,
                status_forcelist=(429, 500, 502, 503, 504),
                allowed_methods=frozenset(["GET", "HEAD"]),
            )
        except TypeError:
            retry = Retry(
                total=4,
                connect=4,
                read=4,
                backoff_factor=0.7,
                status_forcelist=(429, 500, 502, 503, 504),
                method_whitelist=frozenset(["GET", "HEAD"]),  # type: ignore
            )

        adapter = HTTPAdapter(max_retries=retry)
        s.mount("http://", adapter)
        s.mount("https://", adapter)
    except Exception:
        pass

    return s


# -----------------------------
# Candidate model
# -----------------------------


@dataclass(frozen=True)
class Candidate:
    url: str
    source: str
    referer: str = ""


@dataclass
class DownloadResult:
    ok: bool
    filename: str = ""
    source: str = ""
    message: str = ""


@dataclass
class DownloadAttempt:
    ok: bool
    status_code: Optional[int]
    content_type: str
    final_url: str
    followups: List[str]
    message: str = ""


# -----------------------------
# API helpers
# -----------------------------


def _get_json(sess, url: str, timeout_s: int = DEFAULT_TIMEOUT_S) -> Optional[dict]:
    try:
        r = sess.get(url, timeout=timeout_s, headers={"Accept": "application/json"})
        if r.status_code != 200:
            return None
        return r.json()
    except Exception:
        return None


def candidates_from_unpaywall(sess, doi: str, email: str) -> List[Candidate]:
    if not email:
        return []
    doi = normalize_doi(doi)
    api = f"https://api.unpaywall.org/v2/{quote(doi, safe='')}?email={quote(email, safe='')}"
    data = _get_json(sess, api)
    if not isinstance(data, dict):
        return []
    out: List[Candidate] = []

    def add_loc(loc: object, label: str) -> None:
        if not isinstance(loc, dict):
            return
        pdf_url = loc.get("url_for_pdf")
        if isinstance(pdf_url, str) and pdf_url.startswith("http"):
            out.append(Candidate(url=pdf_url, source=f"unpaywall:{label}:url_for_pdf"))
        landing = loc.get("url")
        if isinstance(landing, str) and landing.startswith("http"):
            out.append(Candidate(url=landing, source=f"unpaywall:{label}:landing"))

    add_loc(data.get("best_oa_location"), "best_oa_location")
    for i, loc in enumerate(data.get("oa_locations") or []):
        add_loc(loc, f"oa_locations[{i}]")

    return dedupe_candidates(out)


def candidates_from_openalex(sess, doi: str, mailto: str = "") -> List[Candidate]:
    doi = normalize_doi(doi)
    api = f"https://api.openalex.org/works/https://doi.org/{quote(doi, safe='')}"
    if mailto:
        api += f"?mailto={quote(mailto, safe='')}"
    data = _get_json(sess, api)
    if not isinstance(data, dict):
        return []
    out: List[Candidate] = []

    def harvest_loc(loc: object, label: str) -> None:
        if not isinstance(loc, dict):
            return
        pdf_url = loc.get("pdf_url")
        if isinstance(pdf_url, str) and pdf_url.startswith("http"):
            out.append(Candidate(url=pdf_url, source=f"openalex:{label}:pdf_url"))
        landing = loc.get("landing_page_url")
        if isinstance(landing, str) and landing.startswith("http"):
            out.append(Candidate(url=landing, source=f"openalex:{label}:landing_page_url"))

    harvest_loc(data.get("best_oa_location"), "best_oa_location")
    harvest_loc(data.get("primary_location"), "primary_location")

    locs = data.get("locations")
    if isinstance(locs, list):
        for i, loc in enumerate(locs):
            harvest_loc(loc, f"locations[{i}]")

    oa = data.get("open_access") or {}
    if isinstance(oa, dict):
        oa_url = oa.get("oa_url")
        if isinstance(oa_url, str) and oa_url.startswith("http"):
            out.append(Candidate(url=oa_url, source="openalex:open_access:oa_url"))

    return dedupe_candidates(out)


def candidates_from_semanticscholar(sess, doi: str) -> List[Candidate]:
    doi = normalize_doi(doi)
    api = (
        "https://api.semanticscholar.org/graph/v1/paper/"
        + quote(f"DOI:{doi}", safe="")
        + "?fields=url,openAccessPdf,isOpenAccess"
    )
    data = _get_json(sess, api)
    if not isinstance(data, dict):
        return []
    out: List[Candidate] = []

    oapdf = data.get("openAccessPdf")
    if isinstance(oapdf, dict):
        u = oapdf.get("url")
        if isinstance(u, str) and u.startswith("http"):
            out.append(Candidate(url=u, source="semanticscholar:openAccessPdf:url"))

    landing = data.get("url")
    if isinstance(landing, str) and landing.startswith("http"):
        out.append(Candidate(url=landing, source="semanticscholar:url"))

    return dedupe_candidates(out)


def candidates_from_crossref(sess, doi: str) -> List[Candidate]:
    doi = normalize_doi(doi)
    api = f"https://api.crossref.org/works/{quote(doi, safe='')}"
    data = _get_json(sess, api)
    if not isinstance(data, dict):
        return []
    msg = data.get("message") or {}
    if not isinstance(msg, dict):
        return []
    out: List[Candidate] = []

    for i, link in enumerate(msg.get("link") or []):
        if not isinstance(link, dict):
            continue
        u = link.get("URL")
        if not isinstance(u, str) or not u.startswith("http"):
            continue
        ct = (link.get("content-type") or "").lower()
        if "pdf" in ct or u.lower().endswith(".pdf"):
            out.append(Candidate(url=u, source=f"crossref:link[{i}]:pdf"))
        else:
            out.append(Candidate(url=u, source=f"crossref:link[{i}]:other"))

    landing = msg.get("URL")
    if isinstance(landing, str) and landing.startswith("http"):
        out.append(Candidate(url=landing, source="crossref:URL"))

    res = msg.get("resource") or {}
    if isinstance(res, dict):
        primary = res.get("primary") or {}
        if isinstance(primary, dict):
            u = primary.get("URL")
            if isinstance(u, str) and u.startswith("http"):
                out.append(Candidate(url=u, source="crossref:resource.primary"))

    return dedupe_candidates(out)


def candidates_from_europe_pmc(sess, doi: str) -> List[Candidate]:
    doi = normalize_doi(doi)
    api = f"https://www.ebi.ac.uk/europepmc/webservices/rest/search?query=DOI:{quote(doi, safe='')}&format=json&pageSize=1"
    data = _get_json(sess, api)
    if not isinstance(data, dict):
        return []
    out: List[Candidate] = []

    result_list = data.get("resultList") or {}
    results = []
    if isinstance(result_list, dict):
        results = result_list.get("result") or []
    if not isinstance(results, list) or not results:
        return []
    first = results[0]
    if not isinstance(first, dict):
        return []

    pmcid = first.get("pmcid")
    if isinstance(pmcid, str) and pmcid.startswith("PMC"):
        out.append(Candidate(url=f"https://pmc.ncbi.nlm.nih.gov/articles/{pmcid}/pdf/", source="europepmc:pmcid:pmc-pdf"))
        out.append(Candidate(url=f"https://pmc.ncbi.nlm.nih.gov/articles/{pmcid}/pdf/{pmcid}.pdf", source="europepmc:pmcid:pmc-pdf-direct"))

    ft = first.get("fullTextUrlList") or {}
    if isinstance(ft, dict):
        ft_urls = ft.get("fullTextUrl") or []
        if isinstance(ft_urls, list):
            for item in ft_urls:
                if not isinstance(item, dict):
                    continue
                u = item.get("url")
                if not isinstance(u, str) or not u.startswith("http"):
                    continue
                style = (item.get("documentStyle") or "").lower()
                if "pdf" in style or u.lower().endswith(".pdf"):
                    out.append(Candidate(url=u, source="europepmc:fullTextUrlList:pdf"))
                else:
                    out.append(Candidate(url=u, source="europepmc:fullTextUrlList:landing"))

    return dedupe_candidates(out)


def candidates_from_pmc_oa(sess, doi: str, tool: str = "doi_pdf_downloader_best", email: str = "") -> List[Candidate]:
    """
    NCBI PMC ID Converter (doi -> pmcid) + PMC OA Web Service (pdf links for OA subset).
    """
    doi = normalize_doi(doi)
    if not doi:
        return []
    params = {
        "ids": doi,
        "idtype": "doi",
        "format": "json",
        "tool": tool,
    }
    if email:
        params["email"] = email

    base = "https://pmc.ncbi.nlm.nih.gov/tools/idconv/api/v1/articles/"
    try:
        r = sess.get(base, params=params, timeout=DEFAULT_TIMEOUT_S, headers={"Accept": "application/json"})
        if r.status_code != 200:
            return []
        data = r.json()
    except Exception:
        return []

    records = data.get("records")
    if not isinstance(records, list) or not records:
        return []

    rec0 = records[0]
    if not isinstance(rec0, dict):
        return []
    pmcid = rec0.get("pmcid")
    if not isinstance(pmcid, str) or not pmcid.startswith("PMC"):
        return []
    pmcid = pmcid.strip()

    # OA Web Service (XML)
    oa_url = "https://www.ncbi.nlm.nih.gov/pmc/utils/oa/oa.fcgi"
    try:
        r2 = sess.get(oa_url, params={"id": pmcid}, timeout=DEFAULT_TIMEOUT_S)
        if r2.status_code != 200:
            return []
        root = ET.fromstring(r2.text)
    except Exception:
        return []

    out: List[Candidate] = []
    for link in root.iter("link"):
        fmt = link.attrib.get("format", "").lower()
        href = link.attrib.get("href", "")
        if fmt == "pdf" and href:
            # Convert FTP to HTTPS if needed
            if href.startswith("ftp://ftp.ncbi.nlm.nih.gov/"):
                href = "https://ftp.ncbi.nlm.nih.gov/" + href[len("ftp://ftp.ncbi.nlm.nih.gov/") :]
            out.append(Candidate(url=href, source="pmc:oa_web_service:pdf"))

    return dedupe_candidates(out)


# -----------------------------
# DOI resolution
# -----------------------------


def resolve_doi(sess, doi: str) -> Tuple[Optional[str], bool]:
    doi = normalize_doi(doi)
    if not doi:
        return None, False
    url = f"https://doi.org/{quote(doi, safe='')}"
    try:
        with sess.get(
            url,
            allow_redirects=True,
            timeout=DEFAULT_TIMEOUT_S,
            headers={"Accept": "text/html,application/xhtml+xml,application/pdf;q=0.9,*/*;q=0.8"},
            stream=True,
        ) as r:
            if r.status_code >= 400:
                return None, False
            ctype = (r.headers.get("Content-Type") or "").lower()
            # peek first bytes
            head = b""
            try:
                it = r.iter_content(chunk_size=4096)
                head = next(it, b"")
            except Exception:
                head = b""
            is_pdf = ("pdf" in ctype) or pdf_magic_in(head)
            return r.url, is_pdf
    except Exception:
        return None, False


def doi_accept_pdf_candidate(doi: str) -> Candidate:
    doi = normalize_doi(doi)
    return Candidate(url=f"https://doi.org/{quote(doi, safe='')}", source="doi:accept:application/pdf")


# -----------------------------
# Publisher patterns
# -----------------------------


def extract_elsevier_pii(url: str) -> Optional[str]:
    if not url:
        return None
    m = ELSEVIER_PII_RE.search(url)
    if m:
        return m.group(1)
    return None


def sciencedirect_candidates(doi: str, landing_url: str) -> List[Candidate]:
    """
    Elsevier/ScienceDirect: add Elsevier CDN direct PDF + ScienceDirect endpoints.
    Works best for OA.
    """
    out: List[Candidate] = []
    if not landing_url:
        return out

    pii = extract_elsevier_pii(landing_url)
    if not pii:
        return out

    # Elsevier CDN direct PDFs (often works for OA)
    out.append(Candidate(
        url=f"https://ars.els-cdn.com/content/article/1-s2.0-{pii}-main.pdf",
        source="elsevier:ars.els-cdn:content/article",
        referer=landing_url,
    ))
    # Some Elsevier content uses /content/image/ ... (less common but cheap to try)
    out.append(Candidate(
        url=f"https://ars.els-cdn.com/content/image/1-s2.0-{pii}-main.pdf",
        source="elsevier:ars.els-cdn:content/image",
        referer=landing_url,
    ))

    # ScienceDirect endpoints (may return HTML "Preparing your download")
    base = f"https://www.sciencedirect.com/science/article/pii/{pii}"
    out.append(Candidate(url=f"{base}/pdfft?isDTMRedir=true&download=true", source="sciencedirect:pdfft", referer=landing_url))
    out.append(Candidate(url=f"{base}/pdfft?isDTMRedir=true", source="sciencedirect:pdfft:isDTMRedir", referer=landing_url))
    out.append(Candidate(url=f"{base}/pdfft", source="sciencedirect:pdfft:plain", referer=landing_url))
    out.append(Candidate(url=f"{base}/pdf?isDTMRedir=true&download=true", source="sciencedirect:pdf", referer=landing_url))
    out.append(Candidate(url=base, source="sciencedirect:landing", referer=landing_url))

    return dedupe_candidates(out)


def wiley_candidates(doi: str, landing_url: str = "") -> List[Candidate]:
    doi = normalize_doi(doi)
    if not doi:
        return []
    ref = landing_url or f"https://onlinelibrary.wiley.com/doi/abs/{quote(doi, safe='')}"
    # Wiley frequently uses anti-bot; still worth trying when OA or with legit access.
    return dedupe_candidates([
        Candidate(url=f"https://onlinelibrary.wiley.com/doi/pdf/{quote(doi, safe='')}", source="wiley:doi/pdf", referer=ref),
        Candidate(url=f"https://onlinelibrary.wiley.com/doi/epdf/{quote(doi, safe='')}", source="wiley:doi/epdf", referer=ref),
        Candidate(url=f"https://onlinelibrary.wiley.com/doi/pdfdirect/{quote(doi, safe='')}", source="wiley:doi/pdfdirect", referer=ref),
        Candidate(url=ref, source="wiley:landing", referer=ref),
    ])


def acs_candidates(doi: str, landing_url: str = "") -> List[Candidate]:
    doi = normalize_doi(doi)
    if not doi:
        return []
    ref = landing_url or f"https://pubs.acs.org/doi/{quote(doi, safe='')}"
    return dedupe_candidates([
        Candidate(url=f"https://pubs.acs.org/doi/pdf/{quote(doi, safe='')}", source="acs:doi/pdf", referer=ref),
        Candidate(url=f"https://pubs.acs.org/doi/pdfplus/{quote(doi, safe='')}", source="acs:doi/pdfplus", referer=ref),
        Candidate(url=ref, source="acs:landing", referer=ref),
    ])


def springer_candidates(doi: str, landing_url: str = "") -> List[Candidate]:
    # Springer often exposes citation_pdf_url in HTML; still add typical resolver landing.
    doi = normalize_doi(doi)
    if not doi:
        return []
    ref = landing_url or f"https://link.springer.com/article/{quote(doi, safe='')}"
    return dedupe_candidates([Candidate(url=ref, source="springer:landing", referer=ref)])


def ieee_candidates(doi: str, landing_url: str = "") -> List[Candidate]:
    # IEEE is inconsistent; "stamp.jsp" requires arnumber. We rely mainly on HTML scrape.
    if landing_url:
        return [Candidate(url=landing_url, source="ieee:landing", referer=landing_url)]
    return []


# -----------------------------
# HTML scraping
# -----------------------------


def extract_pdf_like_urls_from_html(html: str, base_url: str) -> List[str]:
    """
    Extract likely PDF URLs from HTML/JS.

    Handles common publisher cases where a 200 OK HTML page embeds the real PDF URL in:
      - <meta name="citation_pdf_url" content="...">
      - <a href="..."> links
      - absolute URLs in scripts/JSON
      - scheme-relative URLs (//host/path)
      - quoted relative URLs inside JS ("/.../pdfft?...")

    Returns a de-duplicated list of absolute URLs.
    """
    if not html:
        return []

    urls: List[str] = []

    # meta citation_pdf_url (can be relative or absolute)
    for m in META_CITATION_PDF_RE.finditer(html):
        u = (m.group(1) or "").strip()
        if not u:
            continue
        u = u.replace("\/", "/").replace("&amp;", "&")
        if u.startswith("//"):
            u = "https:" + u
        if u.startswith("/"):
            u = urljoin(base_url, u)
        if u.startswith("http") and looks_like_pdf_url(u):
            urls.append(u)

    for m in META_PDF_RE.finditer(html):
        u = (m.group(1) or "").strip()
        if not u:
            continue
        u = u.replace("\/", "/").replace("&amp;", "&")
        if u.startswith("//"):
            u = "https:" + u
        if u.startswith("/"):
            u = urljoin(base_url, u)
        if u.startswith("http") and looks_like_pdf_url(u):
            urls.append(u)

    # href links
    for m in HREF_RE.finditer(html):
        href = (m.group(1) or "").strip()
        if not href:
            continue
        href = href.replace("\/", "/").replace("&amp;", "&")
        if href.startswith("//"):
            href = "https:" + href
        if href.startswith("/"):
            href = urljoin(base_url, href)
        if href.startswith("http") and looks_like_pdf_url(href):
            urls.append(href)

    # scheme-relative URLs in text/JS (//host/path)
    for m in DOUBLESLASH_URL_RE.finditer(html):
        u = (m.group(0) or "").strip()
        if not u:
            continue
        u = "https:" + u
        u = u.replace("\/", "/").replace("&amp;", "&")
        if u.startswith("http") and looks_like_pdf_url(u):
            urls.append(u)

    # quoted relative URLs inside JS that look PDF-ish
    for m in REL_PDFISH_RE.finditer(html):
        u = (m.group(1) or "").strip()
        if not u:
            continue
        u = u.replace("\/", "/").replace("&amp;", "&")
        if u.startswith("/"):
            u = urljoin(base_url, u)
        if u.startswith("http") and looks_like_pdf_url(u):
            urls.append(u)

    # absolute URLs in text/JS
    for m in ABS_URL_RE.finditer(html):
        u = (m.group(0) or "").strip().strip(").,;\"'")
        if not u:
            continue
        u = u.replace("\/", "/").replace("&amp;", "&")
        if u.startswith("http") and looks_like_pdf_url(u):
            urls.append(u)

    # de-dupe while preserving order
    seen: Set[str] = set()
    out: List[str] = []
    for u in urls:
        u = (u or "").strip()
        if not u:
            continue
        u = u.split("#", 1)[0]
        if u in seen:
            continue
        seen.add(u)
        out.append(u)

    return out


def html_scrape_candidates(sess, url: str) -> List[Candidate]:
    try:
        r = sess.get(url, timeout=DEFAULT_TIMEOUT_S, headers={"Accept": "text/html,application/xhtml+xml,*/*;q=0.8"})
        if r.status_code != 200:
            return []
        ct = (r.headers.get("Content-Type") or "").lower()
        if "html" not in ct and not ("text/" in ct):
            # If it isn't HTML, still allow the downloader to try it directly.
            return []
        urls = extract_pdf_like_urls_from_html(r.text, base_url=r.url)
        return dedupe_candidates([Candidate(url=u, source="html:scrape", referer=r.url) for u in urls])
    except Exception:
        return []


# -----------------------------
# JabRef integration (optional)
# -----------------------------


def jabref_executable(provided_path: str = "") -> Optional[str]:
    if provided_path:
        p = Path(provided_path).expanduser()
        if p.exists():
            return str(p)
        # allow relying on PATH as given name
        if shutil.which(provided_path):
            return provided_path
    if os.environ.get("JABREF_PATH"):
        p = Path(os.environ["JABREF_PATH"]).expanduser()
        if p.exists():
            return str(p)
    return shutil.which("jabref") or shutil.which("JabRef")


def jabref_fetch_bibtex(doi: str, exe: str) -> Optional[str]:
    doi = normalize_doi(doi)
    if not doi:
        return None
    try:
        # Try a few fetchers depending on JabRef version
        fetchers = ["DOI", "Crossref", "CrossRef", "Medline"]
        for fetcher in fetchers:
            cmd = [exe, "--nogui", f"--fetch={fetcher}:{doi}"]
            proc = subprocess.run(cmd, stdout=subprocess.PIPE, stderr=subprocess.PIPE, text=True, timeout=120)
            if proc.returncode != 0:
                continue
            out = (proc.stdout or "").strip()
            if out.startswith("@"):
                return out
        return None
    except Exception:
        return None


def candidates_from_bibtex_urls(bibtex: str) -> List[Candidate]:
    if not bibtex:
        return []
    out: List[Candidate] = []
    # very simple field parse
    for m in re.finditer(r"\b(url|pdf|file)\s*=\s*\{([^}]+)\}", bibtex, flags=re.IGNORECASE):
        val = m.group(2).strip()
        if val.startswith("http"):
            out.append(Candidate(url=val, source="jabref:bibtex_field:url"))
    return dedupe_candidates(out)


# -----------------------------
# Download + validation
# -----------------------------


def _warmup_referer(sess, referer: str) -> None:
    if not referer:
        return
    try:
        sess.get(referer, timeout=DEFAULT_TIMEOUT_S, headers={"Accept": "text/html,application/xhtml+xml,*/*;q=0.8"})
    except Exception:
        pass


def attempt_download_pdf(
    sess,
    url: str,
    out_path: Path,
    referer: str = "",
    timeout_s: int = DEFAULT_TIMEOUT_S,
    verbose: bool = False,
) -> DownloadAttempt:
    """
    Try to download a URL as PDF. If response is HTML (even with status=200),
    extract follow-up PDF-like URLs and return them.
    """
    tmp = out_path.with_suffix(out_path.suffix + ".part")
    if tmp.exists():
        try:
            tmp.unlink()
        except Exception:
            pass

    headers: Dict[str, str] = {}
    if referer:
        headers["Referer"] = referer

    # Prefer PDF Accept for PDF-ish endpoints
    if looks_like_pdf_url(url):
        headers["Accept"] = "application/pdf,*/*;q=0.8"
    else:
        headers["Accept"] = "text/html,application/xhtml+xml,application/pdf;q=0.9,*/*;q=0.8"

    try:
        with sess.get(url, stream=True, timeout=timeout_s, allow_redirects=True, headers=headers) as r:
            status = r.status_code
            ctype = (r.headers.get("Content-Type") or "").lower()
            final_url = r.url

            # Read first chunk
            it = r.iter_content(chunk_size=1024 * 64)
            first = next(it, b"")
            head = first or b""

            # Success path: looks like PDF
            if status == 200 and (("pdf" in ctype) or pdf_magic_in(head)):
                # stream to disk
                max_bytes = MAX_PDF_MB * 1024 * 1024
                written = 0
                with tmp.open("wb") as f:
                    f.write(first)
                    written += len(first)
                    for chunk in it:
                        if not chunk:
                            continue
                        f.write(chunk)
                        written += len(chunk)
                        if written > max_bytes:
                            raise RuntimeError(f"PDF too large (> {MAX_PDF_MB} MB)")
                tmp.replace(out_path)
                ok = is_pdf_file(out_path)
                if not ok:
                    try:
                        out_path.unlink()
                    except Exception:
                        pass
                return DownloadAttempt(ok=ok, status_code=status, content_type=ctype, final_url=final_url, followups=[], message="downloaded")

            # Not a PDF: sniff more (up to MAX_HTML_SNIFF_BYTES) to find embedded PDF links
            body = head
            remaining = MAX_HTML_SNIFF_BYTES - len(body)
            while remaining > 0:
                try:
                    chunk = next(it, b"")
                except StopIteration:
                    break
                if not chunk:
                    break
                take = chunk[:remaining]
                body += take
                remaining -= len(take)
                if len(take) < len(chunk):
                    break

            followups: List[str] = []
            if body:
                try:
                    text = body.decode("utf-8", errors="ignore")
                    followups = extract_pdf_like_urls_from_html(text, base_url=final_url)
                except Exception:
                    followups = []

            if verbose:
                print(f"    download failed as PDF (status={status}, ctype={ctype}, final={final_url})", flush=True)
                if followups:
                    print(f"    extracted {len(followups)} follow-up URL(s) from HTML", flush=True)

            return DownloadAttempt(
                ok=False,
                status_code=status,
                content_type=ctype,
                final_url=final_url,
                followups=followups,
                message=f"not_pdf_status={status}",
            )
    except Exception as e:
        if verbose:
            print(f"    exception while downloading: {e}", flush=True)
        return DownloadAttempt(ok=False, status_code=None, content_type="", final_url="", followups=[], message=str(e))


def validate_pdf_matches_doi(pdf_path: Path, expected_doi: str, verbose: bool = False) -> Optional[bool]:
    """
    Best-effort validation:
      - if pdf2doi installed and extracts a DOI, check match
      - else if pypdf installed, scan first 2 pages for DOI patterns
    Returns:
      True  => match
      False => mismatch
      None  => unknown
    """
    expected = normalize_doi(expected_doi).lower()
    if not expected:
        return None

    # 1) pdf2doi
    try:
        from pdf2doi import pdf2doi as _pdf2doi  # type: ignore
        try:
            res = _pdf2doi(str(pdf_path))
            if isinstance(res, dict):
                guess = res.get("identifier") or res.get("doi")
            else:
                guess = None
            if isinstance(guess, str) and guess.strip():
                g = normalize_doi(guess).lower()
                if g == expected:
                    return True
                return False
        except Exception:
            pass
    except Exception:
        pass

    # 2) pypdf
    try:
        from pypdf import PdfReader  # type: ignore
        reader = PdfReader(str(pdf_path))
        text = ""
        for page in reader.pages[:2]:
            try:
                text += "\n" + (page.extract_text() or "")
            except Exception:
                continue
        if not text.strip():
            return None
        found = {normalize_doi(m.group(1)).lower() for m in DOI_PATTERN.finditer(text)}
        if not found:
            return None
        if expected in found:
            return True
        return False
    except Exception:
        return None


# -----------------------------
# Candidate collection per DOI
# -----------------------------


def collect_candidates_for_doi(
    sess,
    doi: str,
    unpaywall_email: str = "",
    contact_email: str = "",
    jabref_exe: str = "",
    verbose: bool = False,
) -> Tuple[List[Candidate], List[str]]:
    """
    Returns (candidates, landing_urls)
    """
    doi = normalize_doi(doi)
    candidates: List[Candidate] = []
    landing_urls: List[str] = []

    # Resolve DOI early (gives landing URL and possibly direct PDF)
    landing_url, landing_is_pdf = resolve_doi(sess, doi)
    if landing_url:
        landing_urls.append(landing_url)
        if landing_is_pdf:
            candidates.append(Candidate(url=landing_url, source="doi:resolved:pdf", referer=""))
        else:
            candidates.append(Candidate(url=landing_url, source="doi:resolved:landing", referer=""))

    # OA discovery APIs
    if unpaywall_email:
        candidates.extend(candidates_from_unpaywall(sess, doi, unpaywall_email))
    candidates.extend(candidates_from_openalex(sess, doi, mailto=contact_email or unpaywall_email))
    candidates.extend(candidates_from_semanticscholar(sess, doi))
    candidates.extend(candidates_from_europe_pmc(sess, doi))
    candidates.extend(candidates_from_pmc_oa(sess, doi, email=contact_email))
    candidates.extend(candidates_from_crossref(sess, doi))

    # DOI content negotiation candidate
    candidates.append(doi_accept_pdf_candidate(doi))

    # Accumulate landing URLs from candidates (for later scraping/patterns)
    for c in candidates:
        if not looks_like_pdf_url(c.url):
            landing_urls.append(c.url)

    # JabRef (metadata helper)
    if jabref_exe:
        bib = jabref_fetch_bibtex(doi, jabref_exe)
        if bib:
            candidates.extend(candidates_from_bibtex_urls(bib))

    # Publisher patterns: try to pick the most informative landing/referrer.
    first_landing = landing_urls[0] if landing_urls else ""

    # Elsevier/ScienceDirect: the PII sometimes appears on linkinghub.elsevier.com or other redirects,
    # not necessarily on the first landing URL. Scan all known URLs for a PII.
    all_known_urls = dedupe_strs(landing_urls + [c.url for c in candidates])
    elsevier_ref = ""
    for u in all_known_urls:
        if extract_elsevier_pii(u):
            elsevier_ref = u
            break

    if elsevier_ref:
        candidates.extend(sciencedirect_candidates(doi, elsevier_ref))
    elif first_landing:
        candidates.extend(sciencedirect_candidates(doi, first_landing))

    if first_landing:
        candidates.extend(wiley_candidates(doi, first_landing))
        candidates.extend(acs_candidates(doi, first_landing))
        candidates.extend(springer_candidates(doi, first_landing))
        candidates.extend(ieee_candidates(doi, first_landing))

    # HTML scrape a few landing URLs to discover embedded PDF links
    for land in dedupe_strs(landing_urls)[:6]:
        candidates.extend(html_scrape_candidates(sess, land))

    candidates = dedupe_candidates(candidates)
    candidates.sort(key=candidate_priority)

    if verbose:
        print("  Candidates (top 30):", flush=True)
        for c in candidates[:30]:
            ref = f" (ref={c.referer})" if c.referer else ""
            print(f"    - {c.source}: {c.url}{ref}", flush=True)

    return candidates, landing_urls


def dedupe_strs(xs: Iterable[str]) -> List[str]:
    seen: Set[str] = set()
    out: List[str] = []
    for x in xs:
        x = (x or "").strip()
        if not x:
            continue
        if x in seen:
            continue
        seen.add(x)
        out.append(x)
    return out


def download_pdf_for_doi(
    sess,
    doi: str,
    pdf_dir: Path,
    unpaywall_email: str = "",
    contact_email: str = "",
    jabref_exe: str = "",
    validate: bool = True,
    request_sleep_s: float = DEFAULT_REQ_SLEEP_S,
    request_jitter_s: float = DEFAULT_JITTER_S,
    verbose: bool = False,
) -> DownloadResult:
    doi = normalize_doi(doi)
    if not doi:
        return DownloadResult(ok=False, filename="", source="", message="empty_doi")

    ensure_dir(pdf_dir)
    filename = safe_filename_from_doi(doi)
    out_path = pdf_dir / filename

    if out_path.exists() and is_pdf_file(out_path):
        return DownloadResult(ok=True, filename=filename, source="cache", message="already_exists")

    candidates, _landing_urls = collect_candidates_for_doi(
        sess=sess,
        doi=doi,
        unpaywall_email=unpaywall_email,
        contact_email=contact_email,
        jabref_exe=jabref_exe,
        verbose=verbose,
    )

    queue: List[Candidate] = list(candidates)
    seen_urls: Set[str] = set()
    warmed_referers: Set[str] = set()

    while queue:
        cand = queue.pop(0)
        if cand.url in seen_urls:
            continue
        seen_urls.add(cand.url)

        if verbose:
            print(f"  Trying: {cand.source}: {cand.url}", flush=True)

        # If we have a referer and haven't warmed it up yet, visit it once to get cookies.
        if cand.referer and cand.referer not in warmed_referers:
            _warmup_referer(sess, cand.referer)
            warmed_referers.add(cand.referer)
            sleep_polite(request_sleep_s, request_jitter_s)

        attempt = attempt_download_pdf(sess, cand.url, out_path, referer=cand.referer, verbose=verbose)

        if attempt.ok:
            if not validate:
                return DownloadResult(ok=True, filename=filename, source=cand.source, message="downloaded")

            verdict = validate_pdf_matches_doi(out_path, doi, verbose=verbose)
            if verdict is False:
                # wrong PDF
                if verbose:
                    print("  Validation mismatch; deleting and continuing...", flush=True)
                try:
                    out_path.unlink()
                except Exception:
                    pass
                sleep_polite(request_sleep_s, request_jitter_s)
                continue

            return DownloadResult(ok=True, filename=filename, source=cand.source, message="downloaded_validated" if verdict else "downloaded_unvalidated")

        # 403 handling: warm up with final_url and retry once with same URL (sometimes required)
        if attempt.status_code == 403 and attempt.final_url:
            if attempt.final_url not in warmed_referers:
                _warmup_referer(sess, attempt.final_url)
                warmed_referers.add(attempt.final_url)
                sleep_polite(request_sleep_s, request_jitter_s)
                # retry once
                attempt2 = attempt_download_pdf(sess, cand.url, out_path, referer=cand.referer or attempt.final_url, verbose=verbose)
                if attempt2.ok:
                    if not validate:
                        return DownloadResult(ok=True, filename=filename, source=cand.source, message="downloaded_after_403_warmup")
                    verdict = validate_pdf_matches_doi(out_path, doi, verbose=verbose)
                    if verdict is False:
                        try:
                            out_path.unlink()
                        except Exception:
                            pass
                    else:
                        return DownloadResult(ok=True, filename=filename, source=cand.source, message="downloaded_after_403_warmup")

                # merge followups from retry
                if attempt2.followups:
                    for u in attempt2.followups:
                        if u not in seen_urls:
                            queue.append(Candidate(url=u, source=f"followup:{cand.source}", referer=attempt2.final_url or cand.url))

        # Follow-ups from HTML responses
        if attempt.followups:
            for u in attempt.followups:
                if u not in seen_urls:
                    queue.append(Candidate(url=u, source=f"followup:{cand.source}", referer=attempt.final_url or cand.url))

        sleep_polite(request_sleep_s, request_jitter_s)

    return DownloadResult(ok=False, filename="", source="", message="Undetected")


# -----------------------------
# XLSX processing
# -----------------------------


def find_column(ws, header: str) -> Optional[int]:
    target = header.strip().lower()
    for col in range(1, ws.max_column + 1):
        v = ws.cell(row=1, column=col).value
        if v is None:
            continue
        if str(v).strip().lower() == target:
            return col
    return None


def ensure_pdf_column(ws, doi_col: int) -> int:
    """
    Ensure there is a column titled "PDF".
    Prefer to place it immediately to the right of DOI column (as requested).
    """
    existing = find_column(ws, PDF_COL)
    if existing is not None:
        return existing

    desired = doi_col + 1
    if desired <= ws.max_column:
        # If the adjacent header is empty, reuse it. Otherwise append at end.
        hdr = ws.cell(row=1, column=desired).value
        if hdr is None or str(hdr).strip() == "":
            ws.cell(row=1, column=desired).value = PDF_COL
            return desired

    # Append
    new_col = ws.max_column + 1
    ws.cell(row=1, column=new_col).value = PDF_COL
    return new_col


def process_xlsx(
    input_xlsx: Path,
    output_xlsx: Path,
    pdf_dir: Path,
    unpaywall_email: str,
    contact_email: str,
    jabref_exe: str,
    validate: bool,
    sleep_s: float,
    jitter_s: float,
    request_sleep_s: float,
    request_jitter_s: float,
    verbose: bool,
) -> None:
    import openpyxl  # type: ignore

    wb = openpyxl.load_workbook(str(input_xlsx))
    ws = wb.active

    doi_col = find_column(ws, DOI_COL)
    if doi_col is None:
        raise ValueError(f'Input XLSX must have a column titled "{DOI_COL}" in row 1.')

    pdf_col = ensure_pdf_column(ws, doi_col)

    sess = make_session()

    # Count DOIs
    dois: List[Tuple[int, str]] = []
    for row in range(2, ws.max_row + 1):
        raw = ws.cell(row=row, column=doi_col).value
        doi = normalize_doi(str(raw)) if raw is not None else ""
        if doi:
            dois.append((row, doi))

    total = len(dois)
    for idx, (row, doi) in enumerate(dois, start=1):
        current = ws.cell(row=row, column=pdf_col).value
        # Skip if already downloaded and file exists
        if isinstance(current, str) and current.strip() and current.strip().lower() != "undetected":
            candidate_path = pdf_dir / current.strip()
            if candidate_path.exists() and is_pdf_file(candidate_path):
                print(f"[{idx}/{total}] DOI: {doi}\n  -> already: {current.strip()}", flush=True)
                sleep_polite(sleep_s, jitter_s)
                continue

        print(f"[{idx}/{total}] DOI: {doi}", flush=True)

        res = download_pdf_for_doi(
            sess=sess,
            doi=doi,
            pdf_dir=pdf_dir,
            unpaywall_email=unpaywall_email,
            contact_email=contact_email,
            jabref_exe=jabref_exe,
            validate=validate,
            request_sleep_s=request_sleep_s,
            request_jitter_s=request_jitter_s,
            verbose=verbose,
        )

        if res.ok:
            ws.cell(row=row, column=pdf_col).value = res.filename
            print(f"  -> Downloaded: {res.filename} (via {res.source})", flush=True)
        else:
            ws.cell(row=row, column=pdf_col).value = "Undetected"
            print("  -> Undetected", flush=True)

        wb.save(str(output_xlsx))
        sleep_polite(sleep_s, jitter_s)


# -----------------------------
# Sample XLSX
# -----------------------------


def make_sample_xlsx(path: Path) -> None:
    import openpyxl  # type: ignore

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "DOIs"
    ws.cell(row=1, column=1).value = DOI_COL
    ws.cell(row=2, column=1).value = "10.1016/j.mtbio.2025.101763"
    ws.cell(row=3, column=1).value = "10.1002/adhm.202401358"
    wb.save(str(path))


# -----------------------------
# Main
# -----------------------------


def main(argv: Optional[Sequence[str]] = None) -> int:
    p = argparse.ArgumentParser(description="Download legally available PDFs for DOIs in an XLSX file.")
    p.add_argument("--input", required=False, help='Input XLSX path (must contain column "DOI").')
    p.add_argument("--output", default=None, help="Output XLSX path. Default: overwrite --input.")
    p.add_argument("--pdf-dir", default=None, help='PDF output folder. Default: "./PDFs" next to this script.')
    p.add_argument("--unpaywall-email", default=os.environ.get("UNPAYWALL_EMAIL", ""), help="Email for Unpaywall API (recommended).")
    p.add_argument("--contact-email", default=os.environ.get("CONTACT_EMAIL", ""), help="Contact email for OpenAlex mailto + NCBI.")
    p.add_argument("--jabref", default=os.environ.get("JABREF_PATH", ""), help="Path to JabRef executable (optional).")
    p.add_argument("--no-validate", action="store_true", help="Disable DOI↔PDF validation.")
    p.add_argument("--sleep", type=float, default=DEFAULT_SLEEP_S, help="Base sleep between DOIs (seconds).")
    p.add_argument("--jitter", type=float, default=DEFAULT_JITTER_S, help="Random jitter added to --sleep (seconds).")
    p.add_argument("--request-sleep", type=float, default=DEFAULT_REQ_SLEEP_S, help="Base sleep between URL attempts (seconds).")
    p.add_argument("--request-jitter", type=float, default=DEFAULT_JITTER_S, help="Random jitter added to --request-sleep (seconds).")
    p.add_argument("--verbose", action="store_true", help="Verbose diagnostics (candidates, HTTP failures).")
    p.add_argument("--auto-install", action="store_true", help="Auto-install missing deps (requests, openpyxl).")
    p.add_argument("--make-sample", metavar="PATH", help="Create a sample XLSX at PATH and exit.")
    args = p.parse_args(argv)

    # deps
    if not ensure_import("requests", "requests", auto_install=args.auto_install):
        print("ERROR: requests is not installed. Run: python -m pip install requests", file=sys.stderr)
        return 2
    if not ensure_import("openpyxl", "openpyxl", auto_install=args.auto_install):
        print("ERROR: openpyxl is not installed. Run: python -m pip install openpyxl", file=sys.stderr)
        return 2

    if args.make_sample:
        out = Path(args.make_sample).expanduser().resolve()
        make_sample_xlsx(out)
        print(f"Wrote sample XLSX: {out}")
        return 0

    if not args.input:
        p.print_help()
        return 2

    input_xlsx = Path(args.input).expanduser().resolve()
    if not input_xlsx.exists():
        print(f"ERROR: input XLSX not found: {input_xlsx}", file=sys.stderr)
        return 2

    output_xlsx = Path(args.output).expanduser().resolve() if args.output else input_xlsx

    script_dir = Path(__file__).resolve().parent
    pdf_dir = Path(args.pdf_dir).expanduser().resolve() if args.pdf_dir else (script_dir / "PDFs")
    ensure_dir(pdf_dir)

    jabref_exe = jabref_executable(args.jabref or "") or ""
    if args.verbose and args.jabref and not jabref_exe:
        print(f"[warn] JabRef path not found/usable: {args.jabref}", flush=True)

    validate = not args.no_validate

    # Normalize emails
    unpaywall_email = (args.unpaywall_email or "").strip()
    contact_email = (args.contact_email or "").strip()

    if args.verbose and not unpaywall_email:
        print("[info] UNPAYWALL_EMAIL not set; Unpaywall step will be skipped.", flush=True)

    process_xlsx(
        input_xlsx=input_xlsx,
        output_xlsx=output_xlsx,
        pdf_dir=pdf_dir,
        unpaywall_email=unpaywall_email,
        contact_email=contact_email,
        jabref_exe=jabref_exe,
        validate=validate,
        sleep_s=max(0.0, args.sleep),
        jitter_s=max(0.0, args.jitter),
        request_sleep_s=max(0.0, args.request_sleep),
        request_jitter_s=max(0.0, args.request_jitter),
        verbose=bool(args.verbose),
    )

    print(f"Done. Updated XLSX: {output_xlsx}", flush=True)
    print(f"PDF folder: {pdf_dir}", flush=True)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
