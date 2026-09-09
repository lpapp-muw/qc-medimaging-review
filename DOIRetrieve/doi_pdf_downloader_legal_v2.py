#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
doi_pdf_downloader_legal_v2.py

Input:  XLSX file with a single column titled "DOI" (each row contains a DOI).
Output: The same XLSX is updated with a 2nd column titled "PDF":
        - If a PDF is downloaded: the downloaded PDF filename
        - Otherwise: "Undetected"

PDFs are saved into a folder named "PDFs" located next to this .py file.
If the folder does not exist, it is created.

Important:
- This script uses ONLY legal / open-access endpoints and publisher pages.
- It does NOT use Sci-Hub or any other piracy source.
- For paywalled articles, the result will often be "Undetected" unless your network
  has legitimate access and the publisher allows direct PDF downloads.

Sources / strategies attempted (in roughly this order):
1) OpenAlex (best_oa_location.pdf_url + locations[].pdf_url)  [no API key required]
2) Semantic Scholar (openAccessPdf.url)                       [no API key required]
3) NCBI: PMCID converter + PMC OA Web Service (OA subset PDFs)
4) Unpaywall (if you provide an email)                        [email required]
5) Crossref "link" entries (sometimes includes PDF URLs)
6) DOI resolver (https://doi.org/<doi>) with PDF-leaning headers
7) Publisher URL patterns (ScienceDirect/Wiley) when detected
8) Landing-page HTML scrape (meta citation_pdf_url + hrefs)
9) JabRef CLI (optional) to fetch BibTeX and harvest URLs

Python: tested for syntax on Python 3.8+ (no 3.10+ typing syntax).

Usage:
  python doi_pdf_downloader_legal_v2.py input.xlsx

Optional:
  python doi_pdf_downloader_legal_v2.py input.xlsx --unpaywall-email you@domain.com
  python doi_pdf_downloader_legal_v2.py input.xlsx --contact-email you@domain.com

Create a sample input XLSX (with the two DOIs you provided):
  python doi_pdf_downloader_legal_v2.py --make-sample sample_dois_input.xlsx
"""

from __future__ import annotations

import argparse
import json
import os
import re
import subprocess
import sys
import time
import xml.etree.ElementTree as ET
from dataclasses import dataclass
from pathlib import Path
from typing import Dict, Iterable, List, Optional, Sequence, Set, Tuple
from urllib.parse import quote, urljoin, urlparse

import requests
from openpyxl import Workbook, load_workbook


DEFAULT_TIMEOUT_S = 30
CHUNK_SIZE = 1024 * 64
MAX_PDF_BYTES = 250 * 1024 * 1024  # 250 MB hard stop


DEFAULT_UA = (
    "Mozilla/5.0 (X11; Linux x86_64) "
    "AppleWebKit/537.36 (KHTML, like Gecko) "
    "Chrome/120.0.0.0 Safari/537.36"
)


@dataclass(frozen=True)
class Candidate:
    url: str
    source: str
    referer: Optional[str] = None


@dataclass
class DownloadResult:
    ok: bool
    filename: Optional[str] = None
    source: Optional[str] = None
    message: Optional[str] = None


def normalize_doi(raw: str) -> str:
    """Normalize common DOI presentations into the bare DOI string."""
    s = (raw or "").strip()
    if not s:
        return ""
    s = s.replace("\u200b", "").strip()  # zero-width space
    s_lower = s.lower()

    # Strip common prefixes
    prefixes = [
        "https://doi.org/",
        "http://doi.org/",
        "https://dx.doi.org/",
        "http://dx.doi.org/",
        "doi:",
        "doi ",
    ]
    for p in prefixes:
        if s_lower.startswith(p):
            s = s[len(p) :].strip()
            s_lower = s.lower()
            break

    # Remove surrounding URL noise (rare)
    s = s.strip().strip(".").strip()
    return s


def safe_filename_from_doi(doi: str) -> str:
    """Create a filesystem-safe filename for a DOI."""
    # Replace path separators and other problematic characters.
    base = re.sub(r"[^A-Za-z0-9._-]+", "_", doi)
    base = base.strip("_")
    if not base:
        base = "unknown_doi"
    # Avoid extremely long filenames on some filesystems.
    if len(base) > 180:
        base = base[:180]
    return f"{base}.pdf"


def ensure_dir(p: Path) -> None:
    p.mkdir(parents=True, exist_ok=True)


def make_session() -> requests.Session:
    s = requests.Session()
    s.headers.update(
        {
            "User-Agent": DEFAULT_UA,
            # Prefer PDF, but allow HTML because we sometimes scrape landing pages.
            "Accept": "application/pdf,application/octet-stream;q=0.9,text/html;q=0.8,*/*;q=0.5",
            "Accept-Language": "en-US,en;q=0.9",
        }
    )

    # Best-effort retries; if urllib3 Retry signature differs, we just skip.
    try:
        from requests.adapters import HTTPAdapter
        from urllib3.util.retry import Retry

        try:
            retry = Retry(
                total=4,
                connect=4,
                read=4,
                backoff_factor=0.6,
                status_forcelist=(429, 500, 502, 503, 504),
                allowed_methods=frozenset(["GET", "HEAD"]),
            )
        except TypeError:
            # Older urllib3 used method_whitelist
            retry = Retry(
                total=4,
                connect=4,
                read=4,
                backoff_factor=0.6,
                status_forcelist=(429, 500, 502, 503, 504),
                method_whitelist=frozenset(["GET", "HEAD"]),  # type: ignore
            )

        adapter = HTTPAdapter(max_retries=retry)
        s.mount("http://", adapter)
        s.mount("https://", adapter)
    except Exception:
        pass

    return s


def is_probably_pdf(content_type: str) -> bool:
    ct = (content_type or "").lower()
    return ("application/pdf" in ct) or ("application/octet-stream" in ct)


def download_pdf(
    session: requests.Session,
    url: str,
    dest: Path,
    referer: Optional[str] = None,
    timeout_s: int = DEFAULT_TIMEOUT_S,
) -> bool:
    """Download URL to dest if it is a PDF (by signature)."""
    headers = {}
    if referer:
        headers["Referer"] = referer

    try:
        with session.get(url, headers=headers, stream=True, timeout=timeout_s, allow_redirects=True) as r:
            if r.status_code != 200:
                return False

            # Early abort if huge
            cl = r.headers.get("Content-Length")
            if cl:
                try:
                    if int(cl) > MAX_PDF_BYTES:
                        return False
                except Exception:
                    pass

            # Read first chunk to validate PDF signature
            it = r.iter_content(chunk_size=CHUNK_SIZE)
            try:
                first = next(it)
            except StopIteration:
                return False

            # Some servers may prepend whitespace/newlines.
            if not first.lstrip().startswith(b"%PDF"):
                return False

            # If content-type looks nothing like a PDF, still accept if signature is PDF.
            # Write file
            tmp = dest.with_suffix(dest.suffix + ".part")
            written = 0
            with open(tmp, "wb") as f:
                f.write(first)
                written += len(first)
                for chunk in it:
                    if not chunk:
                        continue
                    f.write(chunk)
                    written += len(chunk)
                    if written > MAX_PDF_BYTES:
                        f.close()
                        try:
                            tmp.unlink()
                        except Exception:
                            pass
                        return False

            tmp.replace(dest)
            return True
    except Exception:
        return False


def dedupe_candidates(cands: Sequence[Candidate]) -> List[Candidate]:
    seen: Set[str] = set()
    out: List[Candidate] = []
    for c in cands:
        u = (c.url or "").strip()
        if not u:
            continue
        if u in seen:
            continue
        seen.add(u)
        out.append(Candidate(u, c.source, c.referer))
    return out


def candidates_from_openalex(session: requests.Session, doi: str, mailto: Optional[str]) -> List[Candidate]:
    """OpenAlex Work endpoint: extract pdf_url and landing_page_url from locations."""
    cands: List[Candidate] = []
    base = f"https://api.openalex.org/works/https://doi.org/{quote(doi, safe='')}"
    if mailto:
        base = f"{base}?mailto={quote(mailto, safe='')}"
    try:
        r = session.get(base, timeout=DEFAULT_TIMEOUT_S, headers={"Accept": "application/json"})
        if r.status_code != 200:
            return []
        data = r.json()
    except Exception:
        return []

    def harvest_location(loc: Dict, label: str) -> None:
        pdf_url = loc.get("pdf_url")
        land = loc.get("landing_page_url")
        if isinstance(pdf_url, str) and pdf_url.strip():
            cands.append(Candidate(pdf_url.strip(), f"openalex:{label}:pdf_url"))
        if isinstance(land, str) and land.strip():
            cands.append(Candidate(land.strip(), f"openalex:{label}:landing"))

    for key in ("best_oa_location", "primary_location"):
        loc = data.get(key)
        if isinstance(loc, dict):
            harvest_location(loc, key)

    locs = data.get("locations")
    if isinstance(locs, list):
        for i, loc in enumerate(locs):
            if isinstance(loc, dict):
                harvest_location(loc, f"locations[{i}]")

    # Some OpenAlex records expose a "doi" field; we ignore.
    return cands


def candidates_from_semanticscholar(session: requests.Session, doi: str) -> List[Candidate]:
    cands: List[Candidate] = []
    encoded_doi = quote(doi, safe="")
    url = (
        f"https://api.semanticscholar.org/graph/v1/paper/DOI:{encoded_doi}"
        f"?fields=openAccessPdf,url,isOpenAccess"
    )
    try:
        r = session.get(url, timeout=DEFAULT_TIMEOUT_S, headers={"Accept": "application/json"})
        if r.status_code != 200:
            return []
        data = r.json()
    except Exception:
        return []

    oapdf = data.get("openAccessPdf")
    if isinstance(oapdf, dict):
        pdf_url = oapdf.get("url")
        if isinstance(pdf_url, str) and pdf_url.strip():
            cands.append(Candidate(pdf_url.strip(), "semanticscholar:openAccessPdf"))

    land = data.get("url")
    if isinstance(land, str) and land.strip():
        cands.append(Candidate(land.strip(), "semanticscholar:landing"))

    return cands


def candidates_from_unpaywall(session: requests.Session, doi: str, email: Optional[str]) -> List[Candidate]:
    if not email:
        return []
    cands: List[Candidate] = []
    url = f"https://api.unpaywall.org/v2/{quote(doi, safe='')}?email={quote(email, safe='')}"
    try:
        r = session.get(url, timeout=DEFAULT_TIMEOUT_S, headers={"Accept": "application/json"})
        if r.status_code != 200:
            return []
        data = r.json()
    except Exception:
        return []

    best = data.get("best_oa_location") or {}
    if isinstance(best, dict):
        pdf = best.get("url_for_pdf") or best.get("url")  # url_for_pdf is preferred
        if isinstance(pdf, str) and pdf.strip():
            cands.append(Candidate(pdf.strip(), "unpaywall:best_oa_location"))
    # Also try other OA locations
    locs = data.get("oa_locations")
    if isinstance(locs, list):
        for i, loc in enumerate(locs):
            if isinstance(loc, dict):
                pdf = loc.get("url_for_pdf") or loc.get("url")
                if isinstance(pdf, str) and pdf.strip():
                    cands.append(Candidate(pdf.strip(), f"unpaywall:oa_locations[{i}]"))

    return cands


def candidates_from_crossref(session: requests.Session, doi: str) -> List[Candidate]:
    cands: List[Candidate] = []
    url = f"https://api.crossref.org/works/{quote(doi, safe='')}"
    try:
        r = session.get(url, timeout=DEFAULT_TIMEOUT_S, headers={"Accept": "application/json"})
        if r.status_code != 200:
            return []
        data = r.json()
    except Exception:
        return []

    msg = data.get("message", {})
    if not isinstance(msg, dict):
        return []

    links = msg.get("link")
    if isinstance(links, list):
        for i, link in enumerate(links):
            if not isinstance(link, dict):
                continue
            u = link.get("URL")
            if not isinstance(u, str) or not u.strip():
                continue
            ct = (link.get("content-type") or "").lower()
            if ("pdf" in ct) or u.lower().endswith(".pdf"):
                cands.append(Candidate(u.strip(), f"crossref:link[{i}]"))

    # Some records have a resource URL (often doi.org) - can still be useful as landing
    res = msg.get("resource")
    if isinstance(res, dict):
        primary = res.get("primary", {})
        if isinstance(primary, dict):
            u = primary.get("URL")
            if isinstance(u, str) and u.strip():
                cands.append(Candidate(u.strip(), "crossref:resource.primary"))

    return cands


def _ftp_to_https(url: str) -> str:
    u = url.strip()
    if u.startswith("ftp://ftp.ncbi.nlm.nih.gov/"):
        return "https://ftp.ncbi.nlm.nih.gov/" + u[len("ftp://ftp.ncbi.nlm.nih.gov/") :]
    return u


def candidates_from_pmc_oa(session: requests.Session, doi: str, tool: str, email: Optional[str]) -> List[Candidate]:
    """
    If the DOI exists in PubMed Central:
      1) Use PMC ID Converter API (doi -> pmcid)
      2) Use PMC OA Web Service to discover direct PDF links (OA subset)
    """
    cands: List[Candidate] = []
    # PMC ID converter (JSON)
    params = {
        "ids": doi,
        "idtype": "doi",
        "format": "json",
        "tool": tool,
    }
    if email:
        params["email"] = email
    base = "https://pmc.ncbi.nlm.nih.gov/tools/idconv/api/v1/articles/"
    try:
        r = session.get(base, params=params, timeout=DEFAULT_TIMEOUT_S, headers={"Accept": "application/json"})
        if r.status_code != 200:
            return []
        data = r.json()
    except Exception:
        return []

    records = data.get("records")
    if not isinstance(records, list) or not records:
        return []

    pmcid = None
    rec0 = records[0]
    if isinstance(rec0, dict):
        pmcid = rec0.get("pmcid")
    if not isinstance(pmcid, str) or not pmcid.strip():
        return []

    pmcid = pmcid.strip()
    # PMC OA Web Service returns XML
    oa_url = "https://www.ncbi.nlm.nih.gov/pmc/utils/oa/oa.fcgi"
    try:
        r2 = session.get(oa_url, params={"id": pmcid}, timeout=DEFAULT_TIMEOUT_S)
        if r2.status_code != 200:
            return []
        xml_text = r2.text
        root = ET.fromstring(xml_text)
    except Exception:
        return []

    # Find <link format="pdf" href="...">
    for link in root.iter("link"):
        fmt = link.attrib.get("format", "").lower()
        href = link.attrib.get("href", "")
        if fmt == "pdf" and href:
            cands.append(Candidate(_ftp_to_https(href), "pmc:oa_web_service", referer=None))

    return cands


def extract_elsevier_pii(url: str) -> Optional[str]:
    u = url or ""
    # /pii/S1234567890123456
    m = re.search(r"/pii/([A-Za-z0-9]+)", u)
    if m:
        return m.group(1)
    # linkinghub.elsevier.com/retrieve/pii/S...
    m = re.search(r"/retrieve/pii/([A-Za-z0-9]+)", u)
    if m:
        return m.group(1)
    return None


def candidates_from_publisher_patterns(landing_url: str, doi: str) -> List[Candidate]:
    """
    Use a few stable publisher URL patterns (legal if OA / you have access).
    """
    cands: List[Candidate] = []
    if not landing_url:
        return cands

    parsed = urlparse(landing_url)
    host = (parsed.netloc or "").lower()

    # Elsevier / ScienceDirect open access often works with /pdfft
    pii = extract_elsevier_pii(landing_url)
    if pii and ("sciencedirect.com" in host or "elsevier.com" in host):
        base = f"https://www.sciencedirect.com/science/article/pii/{pii}/pdfft"
        cands.extend(
            [
                Candidate(base, "pattern:sciencedirect:pdfft", referer=landing_url),
                Candidate(base + "?isDTMRedir=true", "pattern:sciencedirect:pdfft:isDTMRedir", referer=landing_url),
                Candidate(base + "?download=true", "pattern:sciencedirect:pdfft:download", referer=landing_url),
                Candidate(
                    base + "?isDTMRedir=true&download=true",
                    "pattern:sciencedirect:pdfft:isDTMRedir+download",
                    referer=landing_url,
                ),
                # Some OA PDFs may also be on Elsevier CDN (not guaranteed):
                Candidate(f"https://ars.els-cdn.com/content/image/1-s2.0-{pii}-main.pdf", "pattern:elsevier:els-cdn", referer=landing_url),
            ]
        )

    # Wiley (works if OA / your institution has access and Wiley allows it)
    if "onlinelibrary.wiley.com" in host:
        # DOI needs to be URL-encoded in path position (slashes)
        encoded = quote(doi, safe="")
        cands.extend(
            [
                Candidate(f"https://onlinelibrary.wiley.com/doi/pdf/{encoded}", "pattern:wiley:pdf", referer=landing_url),
                Candidate(f"https://onlinelibrary.wiley.com/doi/pdfdirect/{encoded}", "pattern:wiley:pdfdirect", referer=landing_url),
            ]
        )

    return cands


# ---- Simple HTML scraping (meta citation_pdf_url + hrefs) ----
class _LinkExtractor:
    def __init__(self, base_url: str):
        from html.parser import HTMLParser

        class _P(HTMLParser):
            def __init__(self, outer: "_LinkExtractor"):
                super().__init__()
                self.outer = outer

            def handle_starttag(self, tag, attrs):
                a = dict(attrs or [])
                if tag.lower() == "a" and "href" in a:
                    self.outer.hrefs.append(a["href"])
                if tag.lower() == "meta":
                    name = (a.get("name") or a.get("property") or "").lower()
                    content = a.get("content")
                    if content and name:
                        self.outer.metas[name] = content

        self.hrefs: List[str] = []
        self.metas: Dict[str, str] = {}
        self.base_url = base_url
        self._parser = _P(self)

    def feed(self, html: str) -> None:
        self._parser.feed(html or "")


def candidates_from_landing_html(session: requests.Session, landing_url: str) -> List[Candidate]:
    cands: List[Candidate] = []
    if not landing_url:
        return cands
    try:
        r = session.get(landing_url, timeout=DEFAULT_TIMEOUT_S, headers={"Accept": "text/html,*/*;q=0.5"})
        if r.status_code != 200:
            return []
        ct = (r.headers.get("Content-Type") or "").lower()
        if "html" not in ct:
            # If it's already a PDF, that's fine; download stage will handle it.
            cands.append(Candidate(landing_url, "landing:direct_non_html", referer=None))
            return cands
        html = r.text
    except Exception:
        return []

    ex = _LinkExtractor(base_url=landing_url)
    try:
        ex.feed(html)
    except Exception:
        return []

    # Common meta tag used by many publishers
    for key in ("citation_pdf_url", "dc.identifier", "dc.identifier.doi"):
        v = ex.metas.get(key)
        if isinstance(v, str) and v.strip():
            # only treat citation_pdf_url as direct candidate; doi is not pdf
            if key == "citation_pdf_url":
                cands.append(Candidate(v.strip(), f"landing:meta:{key}", referer=landing_url))

    # Heuristic: any href containing ".pdf" or "pdf" in path/query
    for href in ex.hrefs:
        if not isinstance(href, str) or not href.strip():
            continue
        u = href.strip()
        absu = urljoin(landing_url, u)
        low = absu.lower()
        if low.endswith(".pdf") or ".pdf?" in low or "/pdf" in low or "pdfdownload" in low or "pdfft" in low:
            cands.append(Candidate(absu, "landing:href_pdfish", referer=landing_url))

    return cands


def jabref_fetch_bibtex(doi: str, jabref_path: str) -> Optional[str]:
    """
    Use JabRef CLI to fetch a BibTeX entry for a DOI.
    We only use it to harvest URLs (not to download PDFs directly).
    """
    if not jabref_path:
        return None
    try:
        proc = subprocess.run(
            [jabref_path, "--fetch=DOI:" + doi],
            check=False,
            stdout=subprocess.PIPE,
            stderr=subprocess.PIPE,
            text=True,
        )
        if proc.returncode != 0:
            return None
        out = (proc.stdout or "").strip()
        if not out.startswith("@"):
            return None
        return out
    except Exception:
        return None


def candidates_from_bibtex_urls(bibtex: str) -> List[Candidate]:
    """Extract URL-like fields from a BibTeX string."""
    cands: List[Candidate] = []
    if not bibtex:
        return cands
    # Very simple BibTeX field extraction
    for m in re.finditer(r"\b(url|pdf|file)\s*=\s*\{([^}]+)\}", bibtex, flags=re.IGNORECASE):
        val = m.group(2).strip()
        if val.startswith("http://") or val.startswith("https://") or val.startswith("ftp://"):
            cands.append(Candidate(val, "jabref:bibtex_field"))
    return cands


def pdf2doi_validate(pdf_path: Path, expected_doi: str) -> bool:
    """
    Best-effort validation:
    - If pdf2doi is installed and it extracts a DOI that differs from expected, reject.
    - If pdf2doi fails or extracts nothing, accept (to avoid false negatives).
    """
    try:
        from pdf2doi import pdf2doi as _pdf2doi  # type: ignore
    except Exception:
        return True  # cannot validate; accept

    try:
        res = _pdf2doi(str(pdf_path))
        if isinstance(res, dict):
            guess = res.get("identifier") or res.get("doi")
        else:
            guess = None
        if not guess:
            return True
        guess_norm = normalize_doi(str(guess))
        return guess_norm == expected_doi
    except Exception:
        return True


def download_for_doi(
    session: requests.Session,
    doi: str,
    pdf_dir: Path,
    openalex_mailto: Optional[str],
    unpaywall_email: Optional[str],
    contact_email: Optional[str],
    jabref_path: Optional[str],
    validate_with_pdf2doi: bool,
) -> DownloadResult:
    filename = safe_filename_from_doi(doi)
    dest = pdf_dir / filename

    if dest.exists() and dest.stat().st_size > 0:
        return DownloadResult(ok=True, filename=filename, source="cache", message="already downloaded")

    # Collect candidates (ordered by likely OA)
    cands: List[Candidate] = []
    cands.extend(candidates_from_openalex(session, doi, openalex_mailto))
    cands.extend(candidates_from_semanticscholar(session, doi))
    cands.extend(candidates_from_pmc_oa(session, doi, tool="doi_pdf_downloader_legal_v2", email=contact_email))
    cands.extend(candidates_from_unpaywall(session, doi, unpaywall_email))
    cands.extend(candidates_from_crossref(session, doi))

    # DOI resolver as a generic attempt
    doi_url = f"https://doi.org/{doi}"
    cands.append(Candidate(doi_url, "doi.org", referer=None))

    # Collect landing URLs from earlier candidates to drive publisher-pattern and HTML scraping
    landing_urls: List[str] = []
    for c in cands:
        # Heuristic: if it doesn't look like a PDF URL, treat as potential landing page
        low = c.url.lower()
        if not (low.endswith(".pdf") or "pdf" in low):
            landing_urls.append(c.url)

    # Add publisher-pattern candidates + HTML-scraped candidates from landing pages (first few)
    # (Keep this bounded to avoid hammering).
    for land in landing_urls[:6]:
        cands.extend(candidates_from_publisher_patterns(land, doi))
        cands.extend(candidates_from_landing_html(session, land))

    # JabRef optional: fetch BibTeX and harvest URL fields
    if jabref_path:
        bib = jabref_fetch_bibtex(doi, jabref_path=jabref_path)
        if bib:
            cands.extend(candidates_from_bibtex_urls(bib))

    cands = dedupe_candidates(cands)

    # Attempt downloads
    for cand in cands:
        ok = download_pdf(session, cand.url, dest, referer=cand.referer)
        if not ok:
            continue

        # Optional validation
        if validate_with_pdf2doi and not pdf2doi_validate(dest, doi):
            try:
                dest.unlink()
            except Exception:
                pass
            continue

        return DownloadResult(ok=True, filename=filename, source=cand.source)

    # Clean up partials if any
    for part in pdf_dir.glob(filename + ".part"):
        try:
            part.unlink()
        except Exception:
            pass

    return DownloadResult(ok=False, filename=None, source=None, message="Undetected")


def find_column_by_header(ws, header_name: str) -> Optional[int]:
    """Return 1-based column index with matching header (case-insensitive)."""
    header_name = header_name.strip().lower()
    for col in range(1, ws.max_column + 1):
        v = ws.cell(row=1, column=col).value
        if v is None:
            continue
        if str(v).strip().lower() == header_name:
            return col
    return None


def process_xlsx(
    xlsx_path: Path,
    pdf_dir: Path,
    openalex_mailto: Optional[str],
    unpaywall_email: Optional[str],
    contact_email: Optional[str],
    jabref_path: Optional[str],
    validate_with_pdf2doi: bool,
) -> None:
    wb = load_workbook(xlsx_path)
    ws = wb.active

    doi_col = find_column_by_header(ws, "doi")
    if doi_col is None:
        raise ValueError('Input XLSX must have a column with header "DOI" in row 1.')

    # Ensure 2nd column exists and is "PDF" (as requested)
    pdf_col = doi_col + 1
    if ws.max_column < pdf_col:
        ws.cell(row=1, column=pdf_col, value="PDF")
    else:
        header = ws.cell(row=1, column=pdf_col).value
        if header is None or str(header).strip() == "":
            ws.cell(row=1, column=pdf_col, value="PDF")

    session = make_session()

    n = ws.max_row - 1
    for i, row in enumerate(range(2, ws.max_row + 1), start=1):
        raw = ws.cell(row=row, column=doi_col).value
        if raw is None:
            continue
        doi = normalize_doi(str(raw))
        if not doi:
            continue

        print(f"[{i}/{n}] DOI: {doi}")
        res = download_for_doi(
            session=session,
            doi=doi,
            pdf_dir=pdf_dir,
            openalex_mailto=openalex_mailto,
            unpaywall_email=unpaywall_email,
            contact_email=contact_email,
            jabref_path=jabref_path,
            validate_with_pdf2doi=validate_with_pdf2doi,
        )
        if res.ok:
            ws.cell(row=row, column=pdf_col, value=res.filename)
            print(f"  -> Downloaded: {res.filename}  (via {res.source})")
        else:
            ws.cell(row=row, column=pdf_col, value="Undetected")
            print("  -> Undetected")

        # Be polite with public APIs
        time.sleep(0.25)

    wb.save(xlsx_path)
    print(f"\nDone. Updated XLSX: {xlsx_path}")


def make_sample_xlsx(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "DOIs"
    ws["A1"] = "DOI"
    ws["A2"] = "10.1016/j.mtbio.2025.101763"
    ws["A3"] = "10.1002/adhm.202401358"
    wb.save(path)


def main(argv: Optional[Sequence[str]] = None) -> int:
    p = argparse.ArgumentParser(description="Download legally available PDFs for DOIs listed in an XLSX file.")
    p.add_argument("xlsx", nargs="?", help="Path to input XLSX file (will be updated in-place).")
    p.add_argument("--make-sample", metavar="PATH", help="Create a sample XLSX at PATH and exit.")
    p.add_argument(
        "--pdf-dir",
        default=None,
        help='PDF output folder. Default: "PDFs" next to this script.',
    )
    p.add_argument(
        "--unpaywall-email",
        default=os.environ.get("UNPAYWALL_EMAIL"),
        help="Email for Unpaywall API (required by Unpaywall). Can also be set via UNPAYWALL_EMAIL env var.",
    )
    p.add_argument(
        "--contact-email",
        default=os.environ.get("CONTACT_EMAIL"),
        help="Contact email used for NCBI PMC ID converter and OpenAlex mailto. Can also be set via CONTACT_EMAIL env var.",
    )
    p.add_argument(
        "--jabref",
        default=os.environ.get("JABREF_PATH", ""),
        help="Path to JabRef executable for optional BibTeX fetching (e.g., /path/to/JabRef or JabRef.exe).",
    )
    p.add_argument(
        "--validate-with-pdf2doi",
        action="store_true",
        help="If pdf2doi is installed, reject downloads that pdf2doi says belong to a different DOI.",
    )

    args = p.parse_args(argv)

    if args.make_sample:
        out = Path(args.make_sample).expanduser().resolve()
        make_sample_xlsx(out)
        print(f"Wrote sample XLSX: {out}")
        return 0

    if not args.xlsx:
        p.print_help()
        return 2

    xlsx_path = Path(args.xlsx).expanduser().resolve()
    if not xlsx_path.exists():
        print(f"ERROR: XLSX not found: {xlsx_path}", file=sys.stderr)
        return 2

    script_dir = Path(__file__).resolve().parent
    pdf_dir = Path(args.pdf_dir).expanduser().resolve() if args.pdf_dir else (script_dir / "PDFs")
    ensure_dir(pdf_dir)

    contact_email = args.contact_email or None
    openalex_mailto = contact_email  # OpenAlex recommends a mailto param for polite usage
    unpaywall_email = args.unpaywall_email or None

    jabref_path = (args.jabref or "").strip() or None

    process_xlsx(
        xlsx_path=xlsx_path,
        pdf_dir=pdf_dir,
        openalex_mailto=openalex_mailto,
        unpaywall_email=unpaywall_email,
        contact_email=contact_email,
        jabref_path=jabref_path,
        validate_with_pdf2doi=bool(args.validate_with_pdf2doi),
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
