#!/usr/bin/env python3
"""
DOI-centered metadata merge across multiple Excel exports.

Key behaviors:
- Uses DOI.xlsx (default sheet: MERGE) as the reference DOI list.
- Scans ALL other .xlsx files in the same folder (all sheets).
- Matches rows by DOI and extracts common bibliographic fields.
- For each field, selects ONE best value (no concatenation, no "|" aggregation).
- Replaces PublisherOrSource with Journal (journal / source title only).
- Preserves ALL rows from the reference DOI list (no deletions).
- Adds a diagnostic flag (FoundInAnyInputFile) to confirm which DOIs were present in any other input file.

Important implementation detail:
Some exports (e.g., Web of Science) store links as Excel formulas like:
    =HYPERLINK("https%3A%2F%2F...","View Full Record in Web of Science")
Pandas reads only the displayed text. This script extracts the URL from the formula (percent-decoded)
and replaces the displayed text with the real URL for those columns.

Usage:
  python merge_doi_metadata_v2.py --input_dir . --ref DOI.xlsx --ref_sheet MERGE --output merged_DOI_metadata.csv
"""

from __future__ import annotations

import argparse
import csv
import os
import re
from collections import defaultdict
from dataclasses import dataclass
from typing import Dict, List, Optional, Tuple
from urllib.parse import unquote

import numpy as np
import pandas as pd
import openpyxl


DOI_REGEX = re.compile(r"\b10\.\d{4,9}/\S+\b", re.I)
HYPERLINK_RE = re.compile(r'HYPERLINK\("([^"]+)"', re.I)


def normalize_colname(name: object) -> str:
    if name is None:
        return ""
    s = str(name).strip().lower()
    s = re.sub(r"[\s\-_]+", " ", s)
    s = re.sub(r"[^a-z0-9 ]", "", s)
    return s.strip()


def canonicalize_doi(x: object) -> str:
    """Canonical DOI for matching: lowercase, strip doi.org prefix, strip leading 'doi:'."""
    if x is None or (isinstance(x, float) and np.isnan(x)):
        return ""
    s = str(x).strip()
    if not s:
        return ""
    s = re.sub(r"^\s*(https?://(dx\.)?doi\.org/)\s*", "", s, flags=re.I)
    s = re.sub(r"^\s*doi\s*:\s*", "", s, flags=re.I)
    s = s.strip().strip(" \t\r\n.,;)")
    s = s.lstrip("(")
    return s.lower()


def is_doi_like(x: object) -> bool:
    s = canonicalize_doi(x)
    return bool(DOI_REGEX.search(s))


def guess_doi_column(df: pd.DataFrame) -> Optional[str]:
    """
    Heuristically find DOI column:
    - Prefer columns containing 'doi' in header
    - Fall back to any column with high DOI-like ratio
    """
    candidates: List[Tuple[float, int, str]] = []
    for col in df.columns:
        n = normalize_colname(col)
        if "doi" in n:
            vals = df[col].dropna().astype(str).head(5000)
            ratio = (sum(is_doi_like(v) for v in vals) / len(vals)) if len(vals) else 0.0
            pref = 0
            if n == "doi":
                pref = 3
            elif n in {"doi link", "doi url"}:
                pref = 1
            candidates.append((ratio, pref, col))

    if candidates:
        candidates.sort(key=lambda t: (t[0], t[1]), reverse=True)
        if candidates[0][0] > 0.01:
            return candidates[0][2]

    best_col, best_ratio = None, 0.0
    for col in df.columns:
        vals = df[col].dropna().astype(str).head(5000)
        if not len(vals):
            continue
        ratio = sum(is_doi_like(v) for v in vals) / len(vals)
        if ratio > best_ratio:
            best_ratio, best_col = ratio, col

    if best_ratio > 0.05:
        return best_col
    return None


# Field mapping rules (by normalized column name)
FIELD_RULES = {
    "Title": {
        "include": [
            r"^title$",
            r"^article title$",
            r"^document title$",
            r"article title",
            r"document title",
        ],
        "exclude": [r"source title", r"publication title", r"\bjournal\b"],
    },
    "Authors": {
        "include": [
            r"^author full names$",
            r"^authors$",
            r"^author$",
            r"^book authors$",
            r"^book author full names$",
            r"^group authors$",
        ],
        "exclude": [r"keyword", r"affiliat", r"address", r"orcid", r"email", r"funding"],
    },
    "Affiliations": {
        "include": [
            r"^author affiliations$",
            r"^affiliations$",
            r"affiliat",
            r"^addresses$",
            r"reprint addresses",
        ],
        "exclude": [r"email"],
    },
    "PublicationYear": {
        "include": [r"^publication year$", r"^year$", r"published year", r"publication year"],
        "exclude": [r"conference"],
    },
    "Abstract": {
        "include": [r"^abstract$"],
        "exclude": [],
    },
    "Journal": {
        "include": [
            r"^source title$",
            r"source title",
            r"^publication title$",
            r"publication title",
            r"\bjournal\b",
            r"journal title",
            r"container title",
        ],
        "exclude": [r"publisher", r"address", r"city"],
    },
    "DOI": {
        "include": [r"^doi$"],
        "exclude": [r"doi link", r"doi url"],
    },
    "DOILink": {
        "include": [r"doi link", r"doi url"],
        "exclude": [],
    },
    "ManuscriptLink": {
        "include": [r"pdf link", r"full text", r"\blink\b", r"\burl\b", r"record"],
        "exclude": [r"doi link", r"doi url", r"orcid", r"email"],
    },
    "Keywords": {
        "include": [
            r"^author keywords$",
            r"author keywords",
            r"keywords plus",
            r"^keywords$",
            r"ieee terms",
            r"mesh terms",
            r"index terms",
        ],
        "exclude": [r"^authors?$", r"affiliat", r"address"],
    },
}


def _col_matches(n: str, include_pats: List[str], exclude_pats: List[str]) -> bool:
    for pat in exclude_pats:
        if re.search(pat, n):
            return False
    for pat in include_pats:
        if re.search(pat, n):
            return True
    return False


def map_standard_fields(df: pd.DataFrame) -> Dict[str, str]:
    """
    Map each standard field to ONE best-matching column in df.
    Conservative to avoid mixing Authors/Keywords/Affiliations.
    """
    norm_cols = {col: normalize_colname(col) for col in df.columns}
    mapping: Dict[str, str] = {}

    for field, rules in FIELD_RULES.items():
        hits = [col for col, n in norm_cols.items() if _col_matches(n, rules["include"], rules["exclude"])]
        if not hits:
            continue

        def rank(col: str) -> float:
            n = norm_cols[col]
            score = 0.0
            if n.startswith("unnamed"):
                score -= 100.0

            if field == "Title":
                if n == "article title":
                    score += 30
                if n == "document title":
                    score += 25
                if n == "title":
                    score += 10

            if field == "Authors":
                if n == "author full names":
                    score += 40
                if n == "authors":
                    score += 25
                if n == "author":
                    score += 10

            if field == "Affiliations":
                if n == "author affiliations":
                    score += 35
                if n == "affiliations":
                    score += 30
                if n == "addresses":
                    score += 10

            if field == "PublicationYear":
                if n == "publication year":
                    score += 30
                if n == "year":
                    score += 15

            if field == "Journal":
                if n == "source title":
                    score += 40
                if "journal" in n and "publisher" not in n:
                    score += 15

            if field == "Keywords":
                if n == "author keywords":
                    score += 35
                if "keywords plus" in n:
                    score += 20
                if "ieee terms" in n:
                    score += 10

            if field == "ManuscriptLink":
                if "pdf link" in n:
                    score += 30

            if field == "DOILink":
                if n == "doi link":
                    score += 30
                if n == "doi url":
                    score += 25

            score -= len(n) / 60.0
            return score

        best = sorted(set(hits), key=rank, reverse=True)[0]
        mapping[field] = best

    return mapping


def compute_doi_url(doi: str) -> str:
    d = canonicalize_doi(doi)
    return f"https://doi.org/{d}" if d else ""


def _is_nullish(s: str) -> bool:
    return (not s) or (s.strip().lower() in {"nan", "none", "n/a", "na", "null", "-"})

def clean_text(x: object) -> str:
    if x is None:
        return ""
    s = str(x)
    s = s.replace("\r", " ").replace("\n", " ")
    s = re.sub(r"\s+", " ", s).strip()
    if _is_nullish(s):
        return ""
    return s


def text_quality_score(s: str) -> float:
    """
    Score for choosing the single best value among candidates.
    Rewards completeness; penalizes mojibake / replacement chars / control chars.
    """
    s = clean_text(s)
    if not s:
        return 0.0

    length = len(s)
    words = len(re.findall(r"\w+", s))
    unique_chars = len(set(s.lower()))

    replacement = s.count("\ufffd")  # �
    controls = sum(1 for ch in s if ord(ch) < 32 and ch not in "\t\n\r")
    mojibake_hits = len(re.findall(r"[ÃÂâ€žâ€“â€™â€œâ€�]", s))

    weird = replacement * 10 + controls * 5 + mojibake_hits * 2
    weird_ratio = weird / max(1, length)

    score = length + 3 * words + 0.5 * unique_chars
    score -= 200 * weird_ratio
    return score


def is_url(s: str) -> bool:
    return bool(re.match(r"^https?://", s.strip(), flags=re.I))


def link_score(url: str) -> float:
    """
    Prefer likely manuscript/PDF links when available, else keep any URL.
    """
    u = clean_text(url)
    if not u or not is_url(u):
        return 0.0
    ul = u.lower()
    score = 100.0 + len(u)
    if ul.endswith(".pdf"):
        score += 500.0
    if "stamp/stamp.jsp" in ul:
        score += 400.0
    if "pdf" in ul:
        score += 100.0
    return score


def pick_best_value(field: str, values: List[object]) -> str:
    """Return a single best value for a field; do NOT concatenate alternatives."""
    cleaned = [clean_text(v) for v in values if clean_text(v)]
    if not cleaned:
        return ""

    # Deduplicate (case/whitespace insensitive)
    def norm(s: str) -> str:
        return re.sub(r"\s+", " ", s.strip().lower())

    uniq = []
    seen = set()
    for s in cleaned:
        ns = norm(s)
        if ns in seen:
            continue
        uniq.append(s)
        seen.add(ns)

    if not uniq:
        return ""

    if field in {"ManuscriptLink", "DOILink"}:
        # Keep only URL-like values; prefer PDF-ish links
        url_candidates = [s for s in uniq if is_url(s)]
        if not url_candidates:
            return ""
        url_candidates.sort(key=link_score, reverse=True)
        return url_candidates[0]

    uniq.sort(key=text_quality_score, reverse=True)
    return uniq[0]


def pick_year(values: List[object]) -> str:
    """Extract a single best 4-digit year from candidate values (no guessing)."""
    years: List[str] = []
    for v in values:
        s = clean_text(v)
        if not s:
            continue
        years.extend(re.findall(r"(?:19|20)\d{2}", s))
    if not years:
        return ""
    freq: Dict[str, int] = defaultdict(int)
    for y in years:
        freq[y] += 1
    best = sorted(freq.items(), key=lambda t: (t[1], int(t[0])), reverse=True)[0][0]
    return best


def extract_hyperlink_targets(path: str, sheet_name: str, doi_col_name: str, target_col_name: str) -> Dict[str, str]:
    """
    Extract hyperlink targets from Excel formulas like:
      =HYPERLINK("https%3A%2F%2F...","label")
    Returns: {doi_norm: url}
    """
    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=False)
        ws = wb[sheet_name]
    except Exception:
        return {}

    # Header row
    try:
        header_cells = next(ws.iter_rows(min_row=1, max_row=1, values_only=False))
        headers = [c.value for c in header_cells]
    except Exception:
        return {}

    if doi_col_name not in headers or target_col_name not in headers:
        return {}

    doi_idx = headers.index(doi_col_name) + 1
    tgt_idx = headers.index(target_col_name) + 1

    out: Dict[str, str] = {}
    for row in ws.iter_rows(min_row=2, values_only=False):
        doi_val = row[doi_idx - 1].value
        if doi_val is None:
            continue
        doi_norm = canonicalize_doi(doi_val)
        if not doi_norm:
            continue

        val = row[tgt_idx - 1].value
        if isinstance(val, str) and "HYPERLINK" in val.upper():
            m = HYPERLINK_RE.search(val)
            if m:
                out[doi_norm] = unquote(m.group(1))

    return out


def patch_formula_hyperlinks(path: str, sheet_name: str, df: pd.DataFrame, doi_col: str) -> pd.DataFrame:
    """
    Replace displayed text with hyperlink targets for known formula-hyperlink columns if present.
    """
    if df.empty or doi_col not in df.columns:
        return df

    doi_norm = df[doi_col].map(canonicalize_doi)

    # Known columns in your exports that often use HYPERLINK formulas
    candidates = [c for c in ["Web of Science Record", "DOI Link"] if c in df.columns]
    if not candidates:
        return df

    for col in candidates:
        url_map = extract_hyperlink_targets(path, sheet_name, doi_col, col)
        if not url_map:
            continue
        mapped = doi_norm.map(lambda d: url_map.get(d, ""))
        df[col] = np.where(mapped != "", mapped, df[col])

    return df


@dataclass
class SourceTable:
    file: str
    sheet: str
    df: pd.DataFrame
    doi_col: str
    field_map: Dict[str, str]
    index: Dict[str, List[int]]

    @classmethod
    def from_df(cls, file: str, sheet: str, df: pd.DataFrame, doi_col: str, field_map: Dict[str, str]) -> "SourceTable":
        work = df.copy()
        work["_doi_norm"] = work[doi_col].map(canonicalize_doi)
        work = work[work["_doi_norm"] != ""]
        idx: Dict[str, List[int]] = defaultdict(list)
        for i, doi in work["_doi_norm"].items():
            idx[doi].append(i)
        return cls(file=file, sheet=sheet, df=work, doi_col=doi_col, field_map=field_map, index=idx)


def load_sources(input_dir: str, ref_filename: str) -> Tuple[List[SourceTable], List[str]]:
    sources: List[SourceTable] = []
    no_doi_found: List[str] = []

    for fn in sorted(os.listdir(input_dir)):
        if not fn.lower().endswith(".xlsx"):
            continue
        if fn == ref_filename:
            continue

        path = os.path.join(input_dir, fn)
        xl = pd.ExcelFile(path)
        found_any_sheet = False

        for sheet in xl.sheet_names:
            df = pd.read_excel(path, sheet_name=sheet, dtype=str)
            if df.empty:
                continue
            doi_col = guess_doi_column(df)
            if doi_col is None:
                continue

            # Patch formula hyperlinks (WOS export)
            df = patch_formula_hyperlinks(path, sheet, df, doi_col)

            found_any_sheet = True
            fmap = map_standard_fields(df)
            sources.append(SourceTable.from_df(fn, sheet, df, doi_col, fmap))

        if not found_any_sheet:
            no_doi_found.append(fn)

    return sources, no_doi_found


def extract_values(source: SourceTable, doi_norm: str, field: str) -> List[object]:
    col = source.field_map.get(field)
    if not col:
        return []
    idxs = source.index.get(doi_norm, [])
    if not idxs:
        return []
    return [source.df.at[i, col] for i in idxs if col in source.df.columns]


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--input_dir", default=".", help="Folder containing DOI.xlsx and other xlsx files")
    ap.add_argument("--ref", default="DOI.xlsx", help="Reference workbook filename (default: DOI.xlsx)")
    ap.add_argument("--ref_sheet", default="MERGE", help="Sheet inside reference workbook (default: MERGE)")
    ap.add_argument("--output", default="merged_DOI_metadata.csv", help="Output CSV filename")
    args = ap.parse_args()

    ref_path = os.path.join(args.input_dir, args.ref)
    ref_df = pd.read_excel(ref_path, sheet_name=args.ref_sheet, dtype=str)
    if "DOI" not in ref_df.columns:
        raise ValueError(f"Reference sheet must contain a 'DOI' column. Found columns: {list(ref_df.columns)}")

    sources, no_doi_files = load_sources(args.input_dir, args.ref)
    if no_doi_files:
        print("Processed XLSX with NO DOI-like column found (cannot match by DOI):")
        for fn in no_doi_files:
            print(f"  - {fn}")

    if not sources:
        raise RuntimeError("No source sheets with DOI columns found in the other xlsx files.")

    out_fields = [
        "Title",
        "Authors",
        "Affiliations",
        "Journal",
        "PublicationYear",
        "Abstract",
        "Keywords",
        "DOILink",
        "ManuscriptLink",
    ]

    def doi_found_in_any_source(doi_norm: str) -> bool:
        if not doi_norm:
            return False
        return any(doi_norm in s.index for s in sources)

    merged_rows: List[Dict[str, str]] = []
    for _, row in ref_df.iterrows():
        doi_raw = row.get("DOI", "")
        doi_norm = canonicalize_doi(doi_raw)

        out: Dict[str, str] = {col: row[col] for col in ref_df.columns}  # preserve reference columns
        out["Computed_DOI_URL"] = compute_doi_url(doi_raw)

        for field in out_fields:
            candidates: List[object] = []
            for src in sources:
                candidates.extend(extract_values(src, doi_norm, field))
            if field == "PublicationYear":
                out[field] = pick_year(candidates)
            else:
                out[field] = pick_best_value(field, candidates)

        out["FoundInAnyInputFile"] = "YES" if doi_found_in_any_source(doi_norm) else "NO"
        merged_rows.append(out)

    merged_df = pd.DataFrame(merged_rows)

    out_path = os.path.join(args.input_dir, args.output)
    merged_df.to_csv(out_path, sep=";", index=False, encoding="utf-8-sig", quoting=csv.QUOTE_MINIMAL)
    print(f"Wrote: {out_path}")

    # Summary for empty records
    empties = (merged_df[out_fields].fillna("").apply(lambda c: c.astype(str).str.strip() == "").all(axis=1))
    n_empty = int(empties.sum())
    n_total = len(merged_df)
    if n_empty:
        n_empty_not_found = int((empties & (merged_df["FoundInAnyInputFile"] == "NO")).sum())
        print(f"Rows with NO extracted metadata in any field: {n_empty} / {n_total}")
        print(f"  Of these, DOIs not present in ANY other input xlsx: {n_empty_not_found} / {n_empty}")
    else:
        print("All DOIs had at least one extracted metadata field.")


if __name__ == "__main__":
    main()
