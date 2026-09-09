import re
from openpyxl import load_workbook
from collections import Counter, defaultdict

def load(name):
    wb = load_workbook(name, data_only=True); ws = wb['extractions']
    names = {c: ws.cell(2,c).value for c in range(1, ws.max_column+1)}
    data={}
    for r in range(3, ws.max_row+1):
        sn=ws.cell(r,3).value
        if sn in (None,''): continue
        data[sn]={names[c]: ws.cell(r,c).value for c in range(1,ws.max_column+1)}
    wb.close(); return data
hd=load('human.xlsx'); ad=load('ai.xlsx')
common=sorted(set(hd)&set(ad))

CAT=['modality_primary','modality_secondary','imaging_task','anatomy','translational_level','paradigm','hardware_modality','real_or_simulator','simulator_framework','transpilation_level','image_encoding','error_mitigation_type','cross_validation_strategy','baseline_rigour_grade','explainability_mechanism','reporting_standard_adherence']
BOOL=['regulatory_pathway_addressed','prospective_clinical_present','classical_baseline_present','external_test_set','statistical_testing_present','code_release','data_release','weights_release']
EXCLUDED={'dataset_type':'human wrote free-text dataset descriptions, not the enum',
 'dataset_realism_grade':'human used low/moderate/high scale; AI used source-type enum (different construct)',
 'clinical_impact_claim':'human boolean not_reported vs AI extracted claim sentence',
 'computational_cost_training':'human yes/no vs AI free-text cost',
 'computational_cost_inference':'human yes/no vs AI free-text cost',
 'sample_size_justification_reported':'human wrote rationale prose, not boolean (decision 4: exclude)',
 'external_validation_used':'human wrote rationale prose, not boolean (decision 4: exclude)'}

ABSENT={'not reported','nor reported','not_reported','none','none_reported',''}
WHOLE_ALIAS={'radiography':'x-ray','hybrid_quantum_classic':'hybrid_quantum_classical','hybrid_classical_quantum':'hybrid_quantum_classical','hybrid-quantum-classical':'hybrid_quantum_classical','qiskit':'qiskit_aer','hardware-aware':'hardware_aware_mapping','retrospective':'retrospective_clinical','retrospective clinical evaluation':'retrospective_clinical','both?':'both','not reported (simulator)':'simulator'}
TOK_ALIAS={'radiography':'x-ray','qiskit':'qiskit_aer','hardware-aware':'hardware_aware_mapping','hybrid_quantum_classic':'hybrid_quantum_classical','hybrid-quantum-classical':'hybrid_quantum_classical','hybrid_classical_quantum':'hybrid_quantum_classical'}
GRAN={'qcnn':'qml','qnn':'qml','vqc':'hybrid_quantum_classical','qpca':'qml','qica':'qml','hhl':'qml','qsvm':'quantum_kernel','qpso':'quantum_kernel','quantum kernel':'quantum_kernel','quantum annealing':'annealing','qgan':'qml','quantum random walk':'qml','grover':'other','pqc':'hybrid_quantum_classical','pqc (grover)':'other','qsvm (quantum support vector machine)':'quantum_kernel','angle (zzfeaturemap)':'angle','zzfeaturemap':'angle','unary amplitude':'amplitude','brain':'neurology','lung':'pulmonology','lungs':'pulmonology','feature importance':'classical_xai'}

def base(v):
    if v is None: return ''
    return re.sub(r'\s+',' ',str(v).strip().lower())

def to_bool(v):
    s=base(v)
    return 'true' if s.startswith('yes') else 'false'

def ai_bool(v):
    s=base(v)
    return s if s in ('true','false') else to_bool(v)

DR_REQ=['on request','upon request','by request','on demand']
def dr_human(v):
    s=base(v)
    if any(m in s for m in DR_REQ): return 'false'
    return to_bool(v)

def cv_enum(s):
    if s.startswith('not reported') or s in ABSENT: return {'\u2205'}
    if 'fold' in s: return {'k_fold'}
    if 'split' in s or '%' in s or re.search(r'\d+\s*[-/]\s*\d+', s): return {'holdout'}
    return {s}

def norm_set(field,v,lenient):
    s=base(v)
    if field=='cross_validation_strategy': 
        out=cv_enum(s)
        return out
    if s in WHOLE_ALIAS: s=WHOLE_ALIAS[s]
    if s in ABSENT: return {'\u2205'}
    toks=[t.strip() for t in re.split(r'[;,]',s) if t.strip()]
    res=set()
    for t in toks:
        t=TOK_ALIAS.get(t,t)
        if t in ABSENT: t='\u2205'
        if lenient: t=GRAN.get(t,t)
        res.add(t)
    if not res: res={'\u2205'}
    return res

def primary(field,v,lenient):
    s=norm_set(field,v,lenient)
    # deterministic primary: prefer a non-empty token; pick sorted for stability but keep AI-ish
    return sorted(s)[0]

def rel(Hs,As,lenient):
    if not lenient: return Hs==As
    if Hs & As: return True
    if Hs<=As or As<=Hs: return True
    return False

def cohen(pairs):
    # pairs: list of (h,a) single labels
    n=len(pairs)
    if n==0: return None,0.0,0.0
    po=sum(1 for h,a in pairs if h==a)/n
    hc=Counter(h for h,_ in pairs); ac=Counter(a for _,a in pairs)
    cats=set(hc)|set(ac)
    pe=sum((hc[k]/n)*(ac[k]/n) for k in cats)
    if pe>=1.0: return float('nan'),po,pe
    return (po-pe)/(1-pe),po,pe

def build(mode_lenient, onesided):  # onesided in {'exclude','agree','disagree'}; cat fields only
    # returns per-field dict and pooled cells for composite
    perf={}; pooled=[]  # pooled = list of (agree_bool, field, h_label, a_label)
    fields=CAT+BOOL
    for f in fields:
        cells=[]
        for sn in common:
            hv=hd[sn].get(f); av=ad[sn].get(f)
            if base(hv)=='' or base(av)=='': continue
            if f in BOOL:
                h=dr_human(hv) if f=='data_release' else to_bool(hv); a=ai_bool(av)
                cells.append((h==a,f,h,a)); continue
            Hs=norm_set(f,hv,mode_lenient); As=norm_set(f,av,mode_lenient)
            one=(Hs=={'\u2205'})^(As=={'\u2205'})
            if one:
                if onesided=='exclude': continue
                ag = True if onesided=='agree' else False
                # label: for marginal coherence, if agree collapse, else keep primaries
                if ag: lbl=sorted(As if As!={'\u2205'} else Hs)[0]; cells.append((True,f,lbl,lbl))
                else: cells.append((False,f,sorted(Hs)[0],sorted(As)[0]))
                continue
            ag=rel(Hs,As,mode_lenient)
            if ag:
                lbl=sorted(As)[0]; cells.append((True,f,lbl,lbl))
            else:
                cells.append((False,f,sorted(Hs)[0],sorted(As)[0]))
        perf[f]=cells; pooled+=cells
    return perf,pooled

def composite(pooled):
    n=len(pooled)
    po=sum(1 for c in pooled if c[0])/n
    # stratified P_e: cell-weighted mean of per-field marginal chance
    byf=defaultdict(list)
    for ag,f,h,a in pooled: byf[f].append((h,a))
    pe=0.0
    for f,prs in byf.items():
        nf=len(prs); hc=Counter(h for h,_ in prs); ac=Counter(a for _,a in prs)
        cats=set(hc)|set(ac); pef=sum((hc[k]/nf)*(ac[k]/nf) for k in cats)
        pe+=(nf/n)*pef
    k=(po-pe)/(1-pe) if pe<1 else float('nan')
    return k,po,pe,n

print('='*70)
print('FIELD PARTITION')
print('  kappa-eligible categorical:',len(CAT))
print('  kappa-eligible boolean   :',len(BOOL))
print('  EXCLUDED from kappa:')
for f,why in EXCLUDED.items():
    nd=sum(1 for sn in common if base(hd[sn].get(f)) and base(ad[sn].get(f)))
    print('    {:<38} dual={:<3} reason: {}'.format(f,nd,why))

print('\n'+'='*70)
print('PRIMARY: lenient crosswalk, one-sided-NR EXCLUDED, genuine diffs INCLUDED')
perf,pooled=build(True,'exclude')
k,po,pe,n=composite(pooled)
print('  composite stratified kappa = {:.3f}   (P_o={:.3f}  P_e={:.3f}  N={})'.format(k,po,pe,n))
print('\n  per-field [N, observed-agreement, single-field Cohen kappa]:')
print('  {:<36} {:>4} {:>8} {:>9}'.format('field','N','agree','kappa'))
for f in CAT+BOOL:
    prs=[(h,a) for ag,ff,h,a in perf[f]]
    kk,ppo,ppe=cohen(prs)
    ks='deg' if (kk is None or kk!=kk) else '{:.3f}'.format(kk)
    print('  {:<36} {:>4} {:>7.0%} {:>9}'.format(f,len(prs),ppo,ks))

print('\n'+'='*70)
print('SENSITIVITY (composite stratified kappa)')
for lenient,lbl in [(False,'FLOOR alias-only'),(True,'LENIENT')]:
    for onesided in ['exclude','agree','disagree']:
        _,pl=build(lenient,onesided)
        kk,ppo,ppe,nn=composite(pl)
        print('  {:<16} one-sided-NR={:<9} kappa={:.3f}  agree={:.0%}  N={}'.format(lbl,onesided,kk,ppo,nn))

# sequester genuine diffs variant (lenient, exclude one-sided): drop disagreeing cells
_,pl=build(True,'exclude')
kept=[c for c in pl if c[0]]  # only agreeing -> sequester all diffs
kk,ppo,ppe,nn=composite(pl)
seq=[c for c in pl]
# proper sequester: remove genuine-diff cells then recompute
seqcells=[c for c in pl if c[0]]
ks=composite(seqcells) if seqcells else (float('nan'),0,0,0)
print('  {:<16} genuine-diffs SEQUESTERED kappa={:.3f} (degenerate: all-agree by construction)'.format('LENIENT',ks[0]))

# adjudication file: genuine diffs under primary
print('\n'+'='*70)
diffs=[(ff,h,a) for ag,ff,h,a in pooled if not ag]
print('GENUINE categorical disagreements under primary mode:',len(diffs))
byf=Counter(ff for ff,_,_ in diffs)
for f,c in byf.most_common(): print('  {:<36} {}'.format(f,c))

with open('kappa_adjudication_diffs.txt','w') as fh:
    fh.write('GENUINE categorical disagreements (lenient crosswalk, one-sided-NR excluded)\n')
    fh.write('paper\tfield\thuman\tAI\n')
    for sn in common:
        for f in CAT+BOOL:
            hv=hd[sn].get(f); av=ad[sn].get(f)
            if base(hv)=='' or base(av)=='': continue
            if f in BOOL:
                hb=dr_human(hv) if f=='data_release' else to_bool(hv)
                if hb!=ai_bool(av): fh.write('{}\t{}\t{}\t{}\n'.format(sn,f,hv,av))
                continue
            Hs=norm_set(f,hv,True); As=norm_set(f,av,True)
            if (Hs=={'\u2205'})^(As=={'\u2205'}): continue
            if not rel(Hs,As,True): fh.write('{}\t{}\t{}\t{}\n'.format(sn,f,str(hv).replace(chr(10),' '),str(av).replace(chr(10),' ')))
print('\nwrote kappa_adjudication_diffs.txt')

# ---- frozen crosswalk artifact (emitted by the computation; zero drift) ----
import json as _json
_,_plf=build(True,'exclude');  _kf=composite(_plf)
_,_pfl=build(False,'exclude'); _kfl=composite(_pfl)
artifact={
 "_doc":"Frozen crosswalk and method for human-vs-AI extraction inter-rater reliability (Stage 6). Deposited as the OSF artifact backing the section 12 kappa claim. This file plus kappa.py reproduce the reported value.",
 "review":"Quantum Computing for Medical Imaging Applications (IEEE TRPMS)",
 "method":"Composite Cohen's kappa by stratified pooling. Observed agreement P_o pooled across all scored cells. Chance term P_e = cell-count-weighted average of each field's own marginal chance agreement (sum_k p_h(k)*p_a(k) per field). kappa=(P_o-P_e)/(1-P_e). Not a naive pooled confusion matrix.",
 "scope":"Human-filled extraction cells compared against the corresponding AI cell. Cells filled on only one side are coverage gaps and excluded.",
 "locked_parameters":{
   "crosswalk":"frozen as in this file",
   "one_sided_not_reported":"excluded as coverage gap",
   "genuine_categorical_diffs":"included in kappa (not sequestered)",
   "data_release_on_request":"human value containing an on-request marker counts as NOT released (false), overriding any yes prefix",
   "overall_kappa_definition":"composite stratified (this file's method); NOT simple mean of per-field kappa",
   "acceptance_threshold":0.61
 },
 "field_partition":{
   "categorical_scored":CAT,
   "boolean_scored":BOOL,
   "excluded_from_kappa":EXCLUDED
 },
 "normalisation":{
   "absent_family_to_NULL":sorted(ABSENT),
   "whole_string_aliases":WHOLE_ALIAS,
   "token_aliases":TOK_ALIAS,
   "granularity_collapse_lenient":GRAN,
   "boolean_rule":"human: lowercased value startswith 'yes' -> true, else false. AI: literal true/false token.",
   "data_release_on_request_markers":DR_REQ,
   "cross_validation_rule":"contains 'fold' -> k_fold; contains 'split' or '%' or a digit-digit ratio -> holdout; not-reported family -> NULL.",
   "multivalue":"split on ';' or ',', normalise each token, form a set.",
   "matching_relation_lenient":"sets share at least one token OR one set is a subset of the other -> agree."
 },
 "results":{
   "kappa_vocabulary_normalised":round(_kf[0],4),
   "observed_agreement":round(_kf[1],4),
   "P_e":round(_kf[2],4),
   "N_scored_cells":_kf[3],
   "kappa_alias_only_floor":round(_kfl[0],4),
   "floor_observed_agreement":round(_kfl[1],4),
   "interpretation":"Both numbers are reported. The floor isolates synonym/format normalisation only; the difference to the normalised value reflects granularity collapse and subset-overlap scoring. The crosswalk was developed on the observed human-AI pairs (data-derived, not blind) and is disclosed as such."
 }
}
with open('kappa_crosswalk_frozen.json','w') as fh:
    _json.dump(artifact,fh,indent=2,ensure_ascii=False)
print('\nwrote kappa_crosswalk_frozen.json  (kappa_norm={:.4f}  floor={:.4f}  N={})'.format(_kf[0],_kfl[0],_kf[3]))
